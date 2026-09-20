"""SQLite 数据模型：往来单位、单据、行项目、店铺收支/支出日报、供应商结算。"""
import sqlite3
import datetime
import json
import os
import sys

from app import config
from app import settings as settings_mod


def get_conn() -> sqlite3.Connection:
    config.ensure_data_dir()
    conn = sqlite3.connect(config.DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


SCHEMA = """
CREATE TABLE IF NOT EXISTS partners (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    tax_id TEXT DEFAULT '',
    address TEXT DEFAULT '',
    is_supplier INTEGER DEFAULT 0,
    is_customer INTEGER DEFAULT 0
);

CREATE TABLE IF NOT EXISTS documents (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    doc_type TEXT DEFAULT 'FACTURA',
    direction TEXT NOT NULL,             -- compra / venta
    doc_number TEXT DEFAULT '',
    date TEXT,
    partner_id INTEGER,
    partner_name TEXT DEFAULT '',
    tax_id TEXT DEFAULT '',
    base REAL DEFAULT 0,
    iva_rate REAL DEFAULT 0,
    iva_amount REAL DEFAULT 0,
    total REAL DEFAULT 0,
    currency TEXT DEFAULT '',            -- 单据币种 EUR/USD/Bs
    exchange_rate REAL DEFAULT 0,        -- 单据自带汇率（1 USD = X 本国货币）
    cost_total REAL DEFAULT 0,           -- 结转的销售成本
    raw_text TEXT DEFAULT '',
    source_file TEXT DEFAULT '',
    store TEXT DEFAULT '',               -- 入库分店（A店/B店等）
    reviewed INTEGER DEFAULT 0,          -- 审核状态：0=待审核，1=已审核（确认出入库）
    reviewed_at TEXT DEFAULT '',         -- 审核时间
    created_at TEXT DEFAULT (datetime('now','localtime')),
    FOREIGN KEY (partner_id) REFERENCES partners(id)
);

CREATE TABLE IF NOT EXISTS items (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    document_id INTEGER NOT NULL,
    code TEXT DEFAULT '',                -- 商品编码（如 H171）
    description TEXT DEFAULT '',
    qty REAL DEFAULT 1,
    unit_price REAL DEFAULT 0,
    amount REAL DEFAULT 0,
    um TEXT DEFAULT '',                  -- 计量单位（如 STO）
    discount REAL DEFAULT 0,             -- 单项折扣百分比（如 8 表示 8%）
    neto REAL DEFAULT 0,                 -- 折后净额 = amount × (1 - discount/100)
    cost_unit REAL DEFAULT 0,            -- 单位成本（可选）
    FOREIGN KEY (document_id) REFERENCES documents(id) ON DELETE CASCADE
);

CREATE TABLE IF NOT EXISTS daily_income (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    store_a_card REAL DEFAULT 0,
    store_a_ves REAL DEFAULT 0,
    store_a_usd REAL DEFAULT 0,
    store_b_card REAL DEFAULT 0,
    store_b_ves REAL DEFAULT 0,
    store_b_usd REAL DEFAULT 0,
    expense_card REAL DEFAULT 0,
    expense_ves REAL DEFAULT 0,
    expense_usd REAL DEFAULT 0,
    bank_balance REAL DEFAULT 0,
    balance_ves REAL DEFAULT 0,
    balance_usd REAL DEFAULT 0,
    -- 营业额来源拆分：银行卡 / 电子支付 / 现钞（各带币种）
    store_a_card_cur TEXT DEFAULT 'Bs',
    store_a_epay REAL DEFAULT 0,
    store_a_epay_cur TEXT DEFAULT 'USDT',
    store_a_cash REAL DEFAULT 0,
    store_a_cash_cur TEXT DEFAULT 'Bs',
    store_b_card_cur TEXT DEFAULT 'Bs',
    store_b_epay REAL DEFAULT 0,
    store_b_epay_cur TEXT DEFAULT 'USDT',
    store_b_cash REAL DEFAULT 0,
    store_b_cash_cur TEXT DEFAULT 'Bs',
    notes TEXT DEFAULT '',
    created_at TEXT DEFAULT (datetime('now','localtime'))
);

-- 收入日报明细：分店 × 来源（银行卡/电子支付/现钞）× 币种
CREATE TABLE IF NOT EXISTS daily_income_rows (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    income_id INTEGER NOT NULL,
    store TEXT DEFAULT '',               -- 分店名（取自设置的分店列表）
    source TEXT DEFAULT '',              -- 银行卡 / 电子支付 / 现钞
    currency TEXT DEFAULT 'Bs',          -- Bs / USD / CNY / USDT（VES 与 Bs 同币，统一 Bs）
    amount REAL DEFAULT 0,
    notes TEXT DEFAULT '',               -- 本笔备注（每笔独立，不用表头的 notes）
    FOREIGN KEY (income_id) REFERENCES daily_income(id) ON DELETE CASCADE
);

CREATE TABLE IF NOT EXISTS daily_expense (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    summary TEXT DEFAULT '',
    salary REAL DEFAULT 0,
    overtime REAL DEFAULT 0,
    meal REAL DEFAULT 0,
    tax REAL DEFAULT 0,
    utilities REAL DEFAULT 0,
    rent REAL DEFAULT 0,
    municipal REAL DEFAULT 0,
    pay_method TEXT DEFAULT '',          -- 银行卡 / 现金
    pay_currency TEXT DEFAULT '',        -- Bs / 美元 / 人民币
    total_card REAL DEFAULT 0,
    total_ves REAL DEFAULT 0,
    total_usd REAL DEFAULT 0,
    total_cny REAL DEFAULT 0,            -- 人民币合计
    notes TEXT DEFAULT '',
    created_at TEXT DEFAULT (datetime('now','localtime'))
);

-- 支出日报明细：一笔一行（同一天可录入不同币种/付款方式的多笔）
CREATE TABLE IF NOT EXISTS daily_expense_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    summary TEXT DEFAULT '',             -- 摘要
    category TEXT DEFAULT '',            -- 费用类别 key（salary/overtime/...）
    method TEXT DEFAULT '',              -- 银行卡 / 电子支付 / 现金
    currency TEXT DEFAULT '',            -- Bs / 美元 / 人民币 / USDT
    amount REAL DEFAULT 0,
    notes TEXT DEFAULT '',
    created_at TEXT DEFAULT (datetime('now','localtime'))
);

CREATE TABLE IF NOT EXISTS supplier_settlements (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    partner_name TEXT DEFAULT '',
    summary TEXT DEFAULT '',
    doc_number TEXT DEFAULT '',
    supply_amount_ves REAL DEFAULT 0,
    pay_method TEXT DEFAULT '',          -- 银行卡 / 现金
    pay_currency TEXT DEFAULT '',        -- Bs / 美元 / 人民币
    pay_bank REAL DEFAULT 0,             -- 银行转账
    pay_cash_ves REAL DEFAULT 0,         -- 现金（委币）
    pay_cash_usd REAL DEFAULT 0,         -- 现金（美元）
    pay_cash_cny REAL DEFAULT 0,         -- 现金（人民币）
    store TEXT DEFAULT '',               -- 入库分店（A店/B店等）
    notes TEXT DEFAULT '',
    created_at TEXT DEFAULT (datetime('now','localtime'))
);

CREATE TABLE IF NOT EXISTS chart_of_accounts (
    code TEXT PRIMARY KEY,               -- 科目编码（如 1001 / 100201）
    name TEXT NOT NULL,                  -- 科目名称（如 库存现金）
    direction TEXT DEFAULT '借',          -- 方向：借 / 贷
    is_leaf INTEGER DEFAULT 1,           -- 1=明细科目（末级，可录期初）；0=汇总科目
    parent TEXT DEFAULT '',              -- 上级科目编码（按编码前缀推断，顶级为空）
    is_active INTEGER DEFAULT 1
);

CREATE TABLE IF NOT EXISTS opening_balances (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    account_code TEXT NOT NULL,          -- 科目编码
    year TEXT DEFAULT '2026',            -- 会计年度
    opening_balance REAL DEFAULT 0,      -- 年初余额
    cum_debit REAL DEFAULT 0,            -- 本年累计借方
    cum_credit REAL DEFAULT 0,           -- 本年累计贷方
    period_balance REAL DEFAULT 0,       -- 期初余额
    pnl_cum REAL DEFAULT 0,              -- 本年累计损益发生额
    UNIQUE(account_code, year)
);

CREATE TABLE IF NOT EXISTS products (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    code TEXT UNIQUE,                    -- 商品编码 / SKU（与单据行项目 code 对应）
    barcode TEXT DEFAULT '',             -- 条码（扫码枪/小票）
    name TEXT NOT NULL,                  -- 商品名称
    category TEXT DEFAULT '',            -- 分类（食品/饮料/日用品…）
    unit TEXT DEFAULT '',                -- 计量单位 STO/KG/L/UND
    cost_price REAL DEFAULT 0,           -- 成本价（进货价）
    sale_price REAL DEFAULT 0,           -- 售价
    stock_qty REAL DEFAULT 0,            -- 当前库存
    min_stock REAL DEFAULT 0,            -- 库存下限（低于即预警）
    supplier TEXT DEFAULT '',            -- 默认供应商
    notes TEXT DEFAULT ''
);

CREATE TABLE IF NOT EXISTS stock_moves (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    product_id INTEGER NOT NULL,
    date TEXT,                           -- 业务日期（取关联单据日期）
    move_type TEXT,                      -- compra 进货 / venta 出货 / ajuste 调整 / inventario 盘点
    qty REAL,                            -- 变动量：正=入库，负=出库
    balance REAL,                        -- 变动后库存
    ref_doc INTEGER,                     -- 关联单据 id（手工调整留空）
    note TEXT DEFAULT '',
    created_at TEXT DEFAULT (datetime('now','localtime')),
    FOREIGN KEY (product_id) REFERENCES products(id) ON DELETE CASCADE
);
"""


def infer_parent(code: str, codes) -> str:
    """按编码前缀推断上级科目：取最长的、且严格更短的、与 code 同前缀的编码。

    例：100201 → 1002；22210101 → 222101；1001（无更长前缀）→ ''
    """
    best = ""
    for c in codes:
        c = str(c)
        if c == code or len(c) >= len(code):
            continue
        if code.startswith(c) and len(c) > len(best):
            best = c
    return best


def is_leaf_by_tree(code: str, codes) -> int:
    """是否存在下级科目：无下级 → 明细科目(1)，有下级 → 汇总科目(0)。"""
    for c in codes:
        c = str(c)
        if c != code and len(c) > len(code) and c.startswith(code):
            return 0
    return 1


def rebuild_account_tree(conn) -> int:
    """按编码前缀重算全部科目的 parent；仅对尚未设置 is_leaf 的科目补算明细/汇总。

    返回更新条数。
    """
    rows = conn.execute("SELECT code, is_leaf FROM chart_of_accounts").fetchall()
    codes = [r["code"] for r in rows]
    n = 0
    for r in rows:
        parent = infer_parent(r["code"], codes)
        leaf = r["is_leaf"]
        if leaf is None:
            leaf = is_leaf_by_tree(r["code"], codes)
        conn.execute("UPDATE chart_of_accounts SET parent=?, is_leaf=? WHERE code=?",
                     (parent, leaf, r["code"]))
        n += 1
    conn.commit()
    return n


def _migrate_income_rows(conn):
    """把旧版固定 A店/B店 列一次性迁入明细表（仅在明细表为空时调用）。"""
    legacy = (
        ("A店", "store_a_card", "store_a_card_cur", config.INCOME_SOURCE_CARD),
        ("A店", "store_a_epay", "store_a_epay_cur", config.INCOME_SOURCE_EPAY),
        ("A店", "store_a_cash", "store_a_cash_cur", config.INCOME_SOURCE_CASH),
        ("B店", "store_b_card", "store_b_card_cur", config.INCOME_SOURCE_CARD),
        ("B店", "store_b_epay", "store_b_epay_cur", config.INCOME_SOURCE_EPAY),
        ("B店", "store_b_cash", "store_b_cash_cur", config.INCOME_SOURCE_CASH),
    )
    for rec in conn.execute("SELECT * FROM daily_income").fetchall():
        keys = rec.keys()
        for store, amt_col, cur_col, source in legacy:
            if amt_col not in keys:
                continue
            amt = float(rec[amt_col] or 0)
            if not amt:
                continue
            cur = (rec[cur_col] if cur_col in keys else None) or config.INCOME_CUR_BS
            conn.execute(
                """INSERT INTO daily_income_rows
                   (income_id, store, source, currency, amount)
                   VALUES (?,?,?,?,?)""",
                (rec["id"], store, source, cur, amt))


def _migrate_expense_items(conn):
    """把旧版 daily_expense（固定科目列）拆成明细行（仅在明细表为空时调用）。"""
    for rec in conn.execute("SELECT * FROM daily_expense").fetchall():
        keys = rec.keys()
        for key, _label in config.EXPENSE_CATEGORIES:
            if key not in keys:
                continue
            amt = float(rec[key] or 0)
            if not amt:
                continue
            conn.execute(
                """INSERT INTO daily_expense_items
                   (date, summary, category, method, currency, amount, notes)
                   VALUES (?,?,?,?,?,?,?)""",
                (rec["date"], rec["summary"] or "", key,
                 rec["pay_method"] or "", rec["pay_currency"] or "", amt,
                 rec["notes"] or ""))


def _normalize_currency_columns(conn):
    """把历史库里玻利瓦尔的各种写法统一成 Bs（VES / 委内瑞拉玻利瓦尔 / Bs.S …）。

    只需要跑一次（幂等）：改完后再查不到旧值，UPDATE 影响 0 行。
    """
    legacy = ["VES", "VED", "VEF", "BS", "BSS", "BSF", "Bs.S", "BsF",
              "BOLIVAR", "BOLÍVAR", "委内瑞拉玻利瓦尔", "玻利瓦尔"]
    targets = (("documents", "currency"), ("daily_income_rows", "currency"),
               ("daily_expense", "pay_currency"),
               ("daily_expense_items", "currency"),
               ("supplier_settlements", "pay_currency"),
               ("daily_income", "store_a_card_cur"),
               ("daily_income", "store_a_cash_cur"),
               ("daily_income", "store_a_epay_cur"),
               ("daily_income", "store_b_card_cur"),
               ("daily_income", "store_b_cash_cur"),
               ("daily_income", "store_b_epay_cur"))
    for table, col in targets:
        try:
            cols = {r["name"]
                    for r in conn.execute(f"PRAGMA table_info({table})").fetchall()}
        except sqlite3.Error:
            continue
        if col not in cols:
            continue
        ph = ",".join("?" * len(legacy))
        conn.execute(
            f"UPDATE {table} SET {col}=? "
            f"WHERE {col} IS NOT NULL AND TRIM({col}) IN ({ph})",
            [config.CURRENCY_BS] + legacy)
    conn.commit()


def _migrate(conn):
    """旧库补列迁移：documents(currency/exchange_rate/store)、items(code/um/neto/discount)。"""
    doc_cols = {r["name"] for r in conn.execute("PRAGMA table_info(documents)").fetchall()}
    if "currency" not in doc_cols:
        conn.execute("ALTER TABLE documents ADD COLUMN currency TEXT DEFAULT ''")
    if "exchange_rate" not in doc_cols:
        conn.execute("ALTER TABLE documents ADD COLUMN exchange_rate REAL DEFAULT 0")
    if "store" not in doc_cols:
        conn.execute("ALTER TABLE documents ADD COLUMN store TEXT DEFAULT ''")
    if "reviewed" not in doc_cols:
        conn.execute("ALTER TABLE documents ADD COLUMN reviewed INTEGER DEFAULT 0")
        # 存量单据在录入时已联动库存，视为已审核，避免与历史库存重复记账
        conn.execute("UPDATE documents SET reviewed = 1")
    if "reviewed_at" not in doc_cols:
        conn.execute("ALTER TABLE documents ADD COLUMN reviewed_at TEXT DEFAULT ''")

    item_cols = {r["name"] for r in conn.execute("PRAGMA table_info(items)").fetchall()}
    if "code" not in item_cols:
        conn.execute("ALTER TABLE items ADD COLUMN code TEXT DEFAULT ''")
    if "um" not in item_cols:
        conn.execute("ALTER TABLE items ADD COLUMN um TEXT DEFAULT ''")
    if "neto" not in item_cols:
        conn.execute("ALTER TABLE items ADD COLUMN neto REAL DEFAULT 0")
    if "discount" not in item_cols:
        conn.execute("ALTER TABLE items ADD COLUMN discount REAL DEFAULT 0")

    sup_cols = {r["name"]
                for r in conn.execute("PRAGMA table_info(supplier_settlements)").fetchall()}
    if "store" not in sup_cols:
        conn.execute("ALTER TABLE supplier_settlements ADD COLUMN store TEXT DEFAULT ''")
    if "pay_cash_cny" not in sup_cols:
        conn.execute(
            "ALTER TABLE supplier_settlements ADD COLUMN pay_cash_cny REAL DEFAULT 0")

    # 支出日报：补人民币合计列
    exp_cols = {r["name"]
                for r in conn.execute("PRAGMA table_info(daily_expense)").fetchall()}
    if "total_cny" not in exp_cols:
        conn.execute("ALTER TABLE daily_expense ADD COLUMN total_cny REAL DEFAULT 0")

    # 收入日报：营业额来源拆分（银行卡 / 电子支付 / 现钞 + 各自币种）
    inc_cols = {r["name"]
                for r in conn.execute("PRAGMA table_info(daily_income)").fetchall()}
    for col, ddl in (
            ("store_a_card_cur", "TEXT DEFAULT 'Bs'"),
            ("store_a_epay", "REAL DEFAULT 0"),
            ("store_a_epay_cur", "TEXT DEFAULT 'USDT'"),
            ("store_a_cash", "REAL DEFAULT 0"),
            ("store_a_cash_cur", "TEXT DEFAULT 'Bs'"),
            ("store_b_card_cur", "TEXT DEFAULT 'Bs'"),
            ("store_b_epay", "REAL DEFAULT 0"),
            ("store_b_epay_cur", "TEXT DEFAULT 'USDT'"),
            ("store_b_cash", "REAL DEFAULT 0"),
            ("store_b_cash_cur", "TEXT DEFAULT 'Bs'")):
        if col not in inc_cols:
            conn.execute(f"ALTER TABLE daily_income ADD COLUMN {col} {ddl}")

    # 币种归一：VES 与 Bs 是同一种货币，历史库中的 VES / 委内瑞拉玻利瓦尔 统一改为 Bs
    _normalize_currency_columns(conn)

    # 收入日报：把旧版固定 A店/B店 列一次性迁入明细表
    try:
        cnt = conn.execute(
            "SELECT COUNT(*) AS c FROM daily_income_rows").fetchone()["c"]
        if not cnt:
            _migrate_income_rows(conn)
    except sqlite3.OperationalError:
        pass

    # 支出日报：把旧版固定科目列一次性拆入明细表
    try:
        cnt = conn.execute(
            "SELECT COUNT(*) AS c FROM daily_expense_items").fetchone()["c"]
        if not cnt:
            _migrate_expense_items(conn)
    except sqlite3.OperationalError:
        pass

    # 收入日报明细：备注改为每笔独立（旧库把表头备注回填到各明细行，
    # 否则编辑其中一笔的备注会同时改到同一天的其他笔）
    try:
        rcols = {r["name"] for r in conn.execute(
            "PRAGMA table_info(daily_income_rows)").fetchall()}
        if "notes" not in rcols:
            conn.execute(
                "ALTER TABLE daily_income_rows ADD COLUMN notes TEXT DEFAULT ''")
            conn.execute(
                """UPDATE daily_income_rows SET notes = IFNULL((
                       SELECT i.notes FROM daily_income i
                       WHERE i.id = daily_income_rows.income_id), '')""")
    except sqlite3.OperationalError:
        pass

    # 科目表：补 is_leaf（明细科目）/ parent（上级科目）并按编码前缀回填
    acc_cols = {r["name"]
                for r in conn.execute("PRAGMA table_info(chart_of_accounts)").fetchall()}
    need_tree = False
    if "is_leaf" not in acc_cols:
        conn.execute("ALTER TABLE chart_of_accounts ADD COLUMN is_leaf INTEGER")
        need_tree = True
    if "parent" not in acc_cols:
        conn.execute("ALTER TABLE chart_of_accounts ADD COLUMN parent TEXT DEFAULT ''")
        need_tree = True
    if need_tree:
        rebuild_account_tree(conn)


def _load_chart_from_json(conn):
    """首次启动时，若科目表为空且存在 chart_of_accounts.json，则自动导入默认科目。"""
    cnt = conn.execute("SELECT COUNT(*) AS c FROM chart_of_accounts").fetchone()["c"]
    if cnt > 0:
        return 0
    # 优先宿主数据卷（Docker），其次应用目录（本地）
    candidates = [
        os.path.join(config.DATA_DIR, "chart_of_accounts.json"),
        os.path.join(config.APP_DIR, "data", "chart_of_accounts.json"),
    ]
    # PyInstaller 打包（onedir）：随包资源位于 _internal / sys._MEIPASS
    meipass = getattr(sys, "_MEIPASS", "")
    if meipass:
        candidates.append(os.path.join(meipass, "data", "chart_of_accounts.json"))
    candidates.append(os.path.join(config.APP_DIR, "_internal", "data",
                                   "chart_of_accounts.json"))
    path = next((p for p in candidates if os.path.exists(p)), None)
    if path is None:
        return 0
    try:
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
    except Exception:
        return 0
    accounts = data.get("accounts") or []
    n = 0
    for a in accounts:
        code = str(a.get("code", "")).strip()
        name = str(a.get("name", "")).strip()
        if not code or not name:
            continue
        is_leaf = a.get("is_leaf")
        is_leaf = 1 if is_leaf is None else (1 if int(is_leaf) else 0)
        conn.execute(
            """INSERT OR IGNORE INTO chart_of_accounts
               (code, name, direction, is_leaf) VALUES (?,?,?,?)""",
            (code, name, a.get("direction") or "借", is_leaf))
        n += 1
    conn.commit()
    # 全部科目写完后按编码前缀重建层级（parent）
    rebuild_account_tree(conn)
    return n


def init_db():
    conn = get_conn()
    conn.executescript(SCHEMA)
    _migrate(conn)
    _load_chart_from_json(conn)
    conn.commit()
    conn.close()


# ------------------------------------------------------------------ partners

def _find_or_create_partner(conn, name, tax_id, is_supplier, is_customer) -> int:
    if not name:
        return None
    rows = conn.execute(
        "SELECT id FROM partners WHERE name = ? AND tax_id = ?",
        (name, tax_id or "")).fetchall()
    if rows:
        pid = rows[0]["id"]
        conn.execute(
            "UPDATE partners SET is_supplier = MAX(is_supplier, ?), "
            "is_customer = MAX(is_customer, ?) WHERE id = ?",
            (1 if is_supplier else 0, 1 if is_customer else 0, pid))
        return pid
    cur = conn.execute(
        "INSERT INTO partners (name, tax_id, is_supplier, is_customer) VALUES (?,?,?,?)",
        (name, tax_id or "", 1 if is_supplier else 0, 1 if is_customer else 0))
    return cur.lastrowid


def get_partners() -> list:
    conn = get_conn()
    rows = conn.execute("SELECT * FROM partners ORDER BY name").fetchall()
    conn.close()
    return [dict(r) for r in rows]


# ------------------------------------------------------------------ documents

def save_document(doc: dict) -> int:
    """保存单据（含行项目），返回单据 id。doc 结构见 ocr.parser.parse_document。"""
    conn = get_conn()
    try:
        direction = doc.get("direction") or "compra"
        is_sup = 1 if direction == "compra" else 0
        is_cus = 1 if direction == "venta" else 0
        pid = _find_or_create_partner(
            conn, doc.get("partner", "").strip(),
            doc.get("tax_id", ""), is_sup, is_cus)

        d = doc.get("date")
        date_iso = d.strftime("%Y-%m-%d") if isinstance(d, (datetime.date, datetime.datetime)) else d

        cur = conn.execute(
            """INSERT INTO documents
               (doc_type, direction, doc_number, date, partner_id, partner_name,
                tax_id, base, iva_rate, iva_amount, total,
                currency, exchange_rate, cost_total, raw_text, source_file, store)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (doc.get("doc_type", "FACTURA"), direction,
             doc.get("doc_number", ""), date_iso, pid,
             doc.get("partner", ""), doc.get("tax_id", ""),
             doc.get("base", 0.0), doc.get("iva_rate", 0.0),
             doc.get("iva_amount", 0.0), doc.get("total", 0.0),
             config.normalize_currency(doc.get("currency", "")),
             doc.get("exchange_rate", 0.0),
             doc.get("cost_total", 0.0),
             doc.get("raw_text", ""), doc.get("source_file", ""),
             doc.get("store", "")))
        doc_id = cur.lastrowid

        for it in doc.get("items", []) or []:
            conn.execute(
                """INSERT INTO items
                   (document_id, code, description, qty, unit_price, amount,
                    um, discount, neto, cost_unit)
                   VALUES (?,?,?,?,?,?,?,?,?,?)""",
                (doc_id, it.get("code", ""), it.get("desc", ""),
                 it.get("qty", 1), it.get("unit_price", 0),
                 it.get("amount", 0), it.get("um", ""),
                 it.get("discount", 0), it.get("neto", 0),
                 it.get("cost_unit", 0)))
        # 进销存联动：按行项目 code 匹配商品，compra 入库 / venta 出库
        sign = 1 if direction == "compra" else -1
        for it in doc.get("items", []) or []:
            code = (it.get("code") or "").strip()
            if not code:
                continue
            prod = conn.execute(
                "SELECT id FROM products WHERE code = ?", (code,)).fetchone()
            if not prod:
                continue
            qty = float(it.get("qty", 0) or 0) * sign
            new_balance = float(conn.execute(
                "SELECT stock_qty FROM products WHERE id = ?",
                (prod["id"],)).fetchone()["stock_qty"]) + qty
            conn.execute("UPDATE products SET stock_qty = ? WHERE id = ?",
                         (new_balance, prod["id"]))
            conn.execute(
                """INSERT INTO stock_moves
                   (product_id, date, move_type, qty, balance, ref_doc, note)
                   VALUES (?,?,?,?,?,?,?)""",
                (prod["id"], date_iso, direction, qty, new_balance,
                 doc_id, "单据自动记账"))
        conn.commit()
        return doc_id
    finally:
        conn.close()


def list_documents(direction=None, date_from=None, date_to=None, store=None) -> list:
    conn = get_conn()
    sql = "SELECT * FROM documents"
    conds, args = [], []
    if direction:
        conds.append("direction = ?")
        args.append(direction)
    if date_from:
        conds.append("date >= ?")
        args.append(date_from)
    if date_to:
        conds.append("date <= ?")
        args.append(date_to)
    if store:
        conds.append("store = ?")
        args.append(store)
    if conds:
        sql += " WHERE " + " AND ".join(conds)
    sql += " ORDER BY date IS NULL, date, id"
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def list_stores() -> list:
    """单据中出现过的入库分店（去重排序），供筛选/录入下拉复用。"""
    conn = get_conn()
    rows = conn.execute(
        "SELECT DISTINCT store FROM documents WHERE store <> '' ORDER BY store"
    ).fetchall()
    conn.close()
    return [r["store"] for r in rows]


def get_document(doc_id: int) -> dict:
    conn = get_conn()
    row = conn.execute("SELECT * FROM documents WHERE id = ?", (doc_id,)).fetchone()
    doc = dict(row) if row else None
    if doc:
        doc["items"] = [
            {
                "id": r["id"],
                "code": r["code"],
                "desc": r["description"],
                "qty": r["qty"],
                "unit_price": r["unit_price"],
                "amount": r["amount"],
                "um": r["um"],
                "discount": r["discount"],
                "neto": r["neto"],
                "cost_unit": r["cost_unit"],
            }
            for r in conn.execute(
                "SELECT * FROM items WHERE document_id = ?", (doc_id,)).fetchall()
        ]
    conn.close()
    return doc


def _reverse_stock_for_doc(conn, doc_id: int) -> int:
    """撤销某张单据产生的库存影响：回退商品库存并删除其库存流水。

    返回撤销的流水条数。
    """
    moves = conn.execute(
        "SELECT id, product_id, qty FROM stock_moves WHERE ref_doc = ?",
        (doc_id,)).fetchall()
    for m in moves:
        conn.execute("UPDATE products SET stock_qty = stock_qty - ? WHERE id = ?",
                     (float(m["qty"] or 0), m["product_id"]))
    if moves:
        conn.execute("DELETE FROM stock_moves WHERE ref_doc = ?", (doc_id,))
    return len(moves)


def _apply_stock_for_doc(conn, doc_id: int, direction: str, date_iso, items) -> int:
    """按行项目 code 匹配商品登记库存变动：compra 入库(+)，venta 出库(-)。

    返回命中并记账的行数。
    """
    sign = 1 if direction == "compra" else -1
    n = 0
    for it in items or []:
        code = (it.get("code") or "").strip()
        if not code:
            continue
        prod = conn.execute(
            "SELECT id FROM products WHERE code = ?", (code,)).fetchone()
        if not prod:
            continue
        qty = float(it.get("qty") or 0) * sign
        if qty == 0:
            continue
        row = conn.execute(
            "SELECT stock_qty FROM products WHERE id = ?", (prod["id"],)).fetchone()
        new_balance = float(row["stock_qty"] or 0) + qty
        conn.execute("UPDATE products SET stock_qty = ? WHERE id = ?",
                     (new_balance, prod["id"]))
        conn.execute(
            """INSERT INTO stock_moves
               (product_id, date, move_type, qty, balance, ref_doc, note)
               VALUES (?,?,?,?,?,?,?)""",
            (prod["id"], date_iso, direction, qty, new_balance, doc_id,
             "单据自动记账"))
        n += 1
    return n


def update_document(doc_id: int, doc: dict) -> bool:
    """编辑已录入单据：更新主表与行项目，并按新行项目重算库存（保持审核状态不变）。

    doc 结构同 save_document（direction/partner/items 等）；未提供字段沿用原值。
    """
    conn = get_conn()
    try:
        old = conn.execute(
            "SELECT * FROM documents WHERE id = ?", (doc_id,)).fetchone()
        if old is None:
            return False

        direction = doc.get("direction") or old["direction"]
        is_sup = 1 if direction == "compra" else 0
        is_cus = 1 if direction == "venta" else 0
        pid = _find_or_create_partner(
            conn, (doc.get("partner") or old["partner_name"] or "").strip(),
            doc.get("tax_id", old["tax_id"]), is_sup, is_cus)

        d = doc.get("date")
        if isinstance(d, (datetime.date, datetime.datetime)):
            date_iso = d.strftime("%Y-%m-%d")
        else:
            date_iso = d or old["date"]

        conn.execute(
            """UPDATE documents SET
                 doc_type=?, direction=?, doc_number=?, date=?, partner_id=?,
                 partner_name=?, tax_id=?, base=?, iva_rate=?, iva_amount=?,
                 total=?, currency=?, exchange_rate=?, store=?
               WHERE id=?""",
            (doc.get("doc_type", old["doc_type"]), direction,
             doc.get("doc_number", old["doc_number"]), date_iso, pid,
             doc.get("partner", old["partner_name"]),
             doc.get("tax_id", old["tax_id"]),
             doc.get("base", old["base"]), doc.get("iva_rate", old["iva_rate"]),
             doc.get("iva_amount", old["iva_amount"]),
             doc.get("total", old["total"]),
             config.normalize_currency(doc.get("currency", old["currency"])),
             doc.get("exchange_rate", old["exchange_rate"]),
             doc.get("store", old["store"]), doc_id))

        conn.execute("DELETE FROM items WHERE document_id = ?", (doc_id,))
        items = doc.get("items") or []
        for it in items:
            conn.execute(
                """INSERT INTO items
                   (document_id, code, description, qty, unit_price, amount,
                    um, discount, neto, cost_unit)
                   VALUES (?,?,?,?,?,?,?,?,?,?)""",
                (doc_id, it.get("code", ""), it.get("desc", ""),
                 it.get("qty", 1), it.get("unit_price", 0),
                 it.get("amount", 0), it.get("um", ""),
                 it.get("discount", 0), it.get("neto", 0),
                 it.get("cost_unit", 0)))

        # 库存重算：先撤销该单据旧流水，再按编辑后的行项目重新记账
        _reverse_stock_for_doc(conn, doc_id)
        _apply_stock_for_doc(conn, doc_id, direction, date_iso, items)
        conn.commit()
        return True
    finally:
        conn.close()


def set_document_reviewed(doc_id: int, reviewed: bool) -> bool:
    """审核 / 取消审核单据（确认出入库）。

    审核：若该单据尚无库存流水则补记账（幂等），并写入审核时间。
    取消审核：撤销该单据的库存影响，清除审核标记。
    """
    conn = get_conn()
    try:
        row = conn.execute(
            "SELECT * FROM documents WHERE id = ?", (doc_id,)).fetchone()
        if row is None:
            return False
        if reviewed:
            cnt = conn.execute(
                "SELECT COUNT(*) AS c FROM stock_moves WHERE ref_doc = ?",
                (doc_id,)).fetchone()["c"]
            if not cnt:
                items = [dict(r) for r in conn.execute(
                    "SELECT code, qty FROM items WHERE document_id = ?",
                    (doc_id,)).fetchall()]
                _apply_stock_for_doc(conn, doc_id, row["direction"],
                                     row["date"], items)
            conn.execute(
                "UPDATE documents SET reviewed = 1, "
                "reviewed_at = datetime('now','localtime') WHERE id = ?",
                (doc_id,))
        else:
            _reverse_stock_for_doc(conn, doc_id)
            conn.execute(
                "UPDATE documents SET reviewed = 0, reviewed_at = '' WHERE id = ?",
                (doc_id,))
        conn.commit()
        return True
    finally:
        conn.close()


def delete_document(doc_id: int):
    conn = get_conn()
    try:
        # 先回退该单据的库存影响，避免删除后库存虚增/虚减
        _reverse_stock_for_doc(conn, doc_id)
        conn.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
        conn.commit()
    finally:
        conn.close()


# ------------------------------------------------------------------ daily_income

def _date_or_iso(d):
    if isinstance(d, (datetime.date, datetime.datetime)):
        return d.strftime("%Y-%m-%d")
    return d or None


def save_daily_income(rec: dict) -> int:
    """保存收入日报：表头(date/notes) + 明细行(分店 × 来源 × 币种 × 金额)。"""
    conn = get_conn()
    try:
        date_iso = _date_or_iso(rec.get("date"))
        notes = rec.get("notes", "")
        rec_id = rec.get("id")
        if rec_id:
            conn.execute("UPDATE daily_income SET date=?, notes=? WHERE id=?",
                         (date_iso, notes, rec_id))
            conn.execute("DELETE FROM daily_income_rows WHERE income_id=?", (rec_id,))
        else:
            cur = conn.execute(
                "INSERT INTO daily_income (date, notes) VALUES (?,?)",
                (date_iso, notes))
            rec_id = cur.lastrowid
        for row in rec.get("rows", []) or []:
            amt = float(row.get("amount") or 0)
            if not amt:
                continue  # 空行不入库
            # 备注按笔存（row["notes"]）；旧调用方只给表头备注时沿用其值
            row_notes = row["notes"] if "notes" in row else notes
            conn.execute(
                """INSERT INTO daily_income_rows
                   (income_id, store, source, currency, amount, notes)
                   VALUES (?,?,?,?,?,?)""",
                (rec_id, row.get("store", ""), row.get("source", ""),
                 config.normalize_currency(
                     row.get("currency") or config.INCOME_CUR_BS), amt,
                 row_notes or ""))
        conn.commit()
        return rec_id
    finally:
        conn.close()


def list_daily_income(date_from=None, date_to=None) -> list:
    conn = get_conn()
    sql = "SELECT * FROM daily_income"
    conds, args = [], []
    if date_from:
        conds.append("date >= ?")
        args.append(date_from)
    if date_to:
        conds.append("date <= ?")
        args.append(date_to)
    if conds:
        sql += " WHERE " + " AND ".join(conds)
    sql += " ORDER BY date IS NULL, date, id"
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_daily_income(rec_id: int) -> dict:
    """返回日报表头 + 明细行（rows）。"""
    conn = get_conn()
    try:
        row = conn.execute("SELECT * FROM daily_income WHERE id = ?",
                           (rec_id,)).fetchone()
        if not row:
            return None
        rec = dict(row)
        rec["rows"] = [dict(r) for r in conn.execute(
            "SELECT * FROM daily_income_rows WHERE income_id = ? ORDER BY id",
            (rec_id,)).fetchall()]
        return rec
    finally:
        conn.close()


def list_daily_income_full(date_from=None, date_to=None) -> list:
    """返回带明细行 rows 的日报列表。"""
    recs = list_daily_income(date_from, date_to)
    if not recs:
        return []
    conn = get_conn()
    try:
        ids = [r["id"] for r in recs]
        ph = ",".join("?" * len(ids))
        by_id = {}
        for row in conn.execute(
                f"SELECT * FROM daily_income_rows WHERE income_id IN ({ph}) "
                "ORDER BY id", ids).fetchall():
            by_id.setdefault(row["income_id"], []).append(dict(row))
        for r in recs:
            r["rows"] = by_id.get(r["id"], [])
        return recs
    finally:
        conn.close()


def delete_daily_income(rec_id: int):
    conn = get_conn()
    try:
        conn.execute("DELETE FROM daily_income_rows WHERE income_id = ?", (rec_id,))
        conn.execute("DELETE FROM daily_income WHERE id = ?", (rec_id,))
        conn.commit()
    finally:
        conn.close()


def get_or_create_daily_income(date, notes: str = "") -> int:
    """按日期取当日日报表头 id，没有则新建（用于追加明细行）。"""
    d = _date_or_iso(date)
    conn = get_conn()
    try:
        row = conn.execute(
            "SELECT id FROM daily_income WHERE date = ? ORDER BY id DESC LIMIT 1",
            (d,)).fetchone()
        if row:
            return row["id"]
        cur = conn.execute("INSERT INTO daily_income (date, notes) VALUES (?,?)",
                           (d, notes))
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def update_daily_income_header(income_id: int, date, notes: str = ""):
    """只更新日报表头（日期 / 备注），不影响已有明细行。"""
    conn = get_conn()
    try:
        conn.execute("UPDATE daily_income SET date=?, notes=? WHERE id=?",
                     (_date_or_iso(date), notes, income_id))
        conn.commit()
    finally:
        conn.close()


def add_daily_income_row(income_id: int, row: dict) -> int:
    """在指定日报下追加一条明细行（分店 × 支付方式 × 币种 × 金额 + 本笔备注）。"""
    conn = get_conn()
    try:
        cur = conn.execute(
            """INSERT INTO daily_income_rows
               (income_id, store, source, currency, amount, notes)
               VALUES (?,?,?,?,?,?)""",
            (income_id, (row.get("store") or "").strip(), row.get("source", ""),
             config.normalize_currency(row.get("currency") or config.INCOME_CUR_BS),
             float(row.get("amount") or 0), row.get("notes") or ""))
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


def update_daily_income_row(row_id: int, row: dict) -> int:
    """更新一条收入明细行（含本笔备注），返回其所属日报 id（行不存在返回 0）。"""
    conn = get_conn()
    try:
        cur = conn.execute("SELECT income_id FROM daily_income_rows WHERE id = ?",
                           (row_id,)).fetchone()
        if not cur:
            return 0
        conn.execute(
            """UPDATE daily_income_rows
               SET store=?, source=?, currency=?, amount=?, notes=? WHERE id=?""",
            ((row.get("store") or "").strip(), row.get("source", ""),
             config.normalize_currency(row.get("currency") or config.INCOME_CUR_BS),
             float(row.get("amount") or 0), row.get("notes") or "", row_id))
        conn.commit()
        return cur["income_id"]
    finally:
        conn.close()


def move_daily_income_row(row_id: int, income_id: int) -> int:
    """把一条收入明细行改挂到另一个日报表头（改日期时用）。

    原表头若已无任何明细行则一并删除。
    """
    conn = get_conn()
    try:
        old = conn.execute("SELECT income_id FROM daily_income_rows WHERE id = ?",
                           (row_id,)).fetchone()
        conn.execute("UPDATE daily_income_rows SET income_id=? WHERE id=?",
                     (income_id, row_id))
        if old and old["income_id"] != income_id:
            left = conn.execute(
                "SELECT COUNT(*) AS c FROM daily_income_rows WHERE income_id = ?",
                (old["income_id"],)).fetchone()["c"]
            if not left:
                conn.execute("DELETE FROM daily_income WHERE id = ?",
                             (old["income_id"],))
        conn.commit()
        return income_id
    finally:
        conn.close()


def get_daily_income_row(row_id: int) -> dict:
    """按明细行 id 取一笔收入（日期取所属日报，备注取本笔），供再编辑使用。"""
    conn = get_conn()
    try:
        row = conn.execute(
            """SELECT r.id AS row_id, r.income_id, r.store, r.source, r.currency,
                      r.amount,
                      IFNULL(r.notes, i.notes) AS notes,
                      i.date AS date
               FROM daily_income_rows r
               JOIN daily_income i ON i.id = r.income_id
               WHERE r.id = ?""", (row_id,)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def _row_notes(row: dict, rec: dict) -> str:
    """明细行备注：以本笔为准；旧库/异常数据回退到表头备注。"""
    val = row.get("notes")
    if val is None:
        return rec.get("notes") or ""
    return val


def list_daily_income_rows(date_from=None, date_to=None, limit: int = None) -> list:
    """收入日报展开成一笔一行（备注为每笔独立），供列表展示与再编辑。"""
    out = []
    for rec in list_daily_income_full(date_from, date_to):
        for row in rec.get("rows", []) or []:
            out.append({
                "row_id": row.get("id"),
                "income_id": rec.get("id"),
                "date": rec.get("date") or "",
                "notes": _row_notes(row, rec),
                "store": row.get("store") or "",
                "source": row.get("source") or "",
                "currency": row.get("currency") or "",
                "amount": float(row.get("amount") or 0),
            })
    return out[-limit:] if limit else out


def delete_daily_income_row(row_id: int):
    """只删除一条明细行；父记录若无剩余行则一并删除。"""
    conn = get_conn()
    try:
        row = conn.execute(
            "SELECT income_id FROM daily_income_rows WHERE id = ?",
            (row_id,)).fetchone()
        conn.execute("DELETE FROM daily_income_rows WHERE id = ?", (row_id,))
        if row:
            left = conn.execute(
                "SELECT COUNT(*) AS c FROM daily_income_rows WHERE income_id = ?",
                (row["income_id"],)).fetchone()["c"]
            if not left:
                conn.execute("DELETE FROM daily_income WHERE id = ?",
                             (row["income_id"],))
        conn.commit()
    finally:
        conn.close()


def daily_income_by_account(date_from=None, date_to=None) -> dict:
    """把收入日报明细按「支付方式 + 币种」归集到报表科目 key。

    cash=库存现金 / bank=银行存款 / crypto=其他货币资金（扫码支付）；
    稳定币（USDT）不记现金，现钞来源的稳定币改记银行存款。
    返回 {"cash": {币种: 金额}, ...}；金额为原始币种，换算由调用方处理。
    """
    out = {"cash": {}, "bank": {}, "crypto": {}}
    for r in list_daily_income_full(date_from, date_to):
        for row in r.get("rows", []):
            amt = float(row.get("amount") or 0)
            if not amt:
                continue
            cur = config.normalize_currency(row.get("currency")) or config.INCOME_CUR_BS
            key = config.income_account_key(row.get("source"), cur)
            out[key][cur] = out[key].get(cur, 0.0) + amt
    return out


# ------------------------------------------------------------------ daily_expense

def _calc_expense_totals(rec: dict):
    """根据方式/币种把各科目合计归集到 total_card/total_ves/total_usd/total_cny。"""
    total = sum(float(rec.get(k, 0) or 0) for k, _ in config.EXPENSE_CATEGORIES)
    method = rec.get("pay_method", "")
    currency = config.normalize_currency(rec.get("pay_currency", ""))
    card = total if method == config.PAY_METHOD_CARD else 0
    cash = method == config.PAY_METHOD_CASH
    ves = total if (cash and currency == config.PAY_CURRENCY_BS) else 0
    usd = total if (cash and currency == config.PAY_CURRENCY_USD) else 0
    cny = total if (cash and currency == config.PAY_CURRENCY_CNY) else 0
    return card, ves, usd, cny


def save_daily_expense(rec: dict) -> int:
    conn = get_conn()
    try:
        card, ves, usd, cny = _calc_expense_totals(rec)
        data = {
            "date": _date_or_iso(rec.get("date")),
            "summary": rec.get("summary", ""),
            "salary": rec.get("salary", 0),
            "overtime": rec.get("overtime", 0),
            "meal": rec.get("meal", 0),
            "tax": rec.get("tax", 0),
            "utilities": rec.get("utilities", 0),
            "rent": rec.get("rent", 0),
            "municipal": rec.get("municipal", 0),
            "pay_method": rec.get("pay_method", ""),
            "pay_currency": config.normalize_currency(rec.get("pay_currency", "")),
            "total_card": card,
            "total_ves": ves,
            "total_usd": usd,
            "total_cny": cny,
            "notes": rec.get("notes", ""),
        }
        rec_id = rec.get("id")
        if rec_id:
            conn.execute(
                """UPDATE daily_expense SET
                   date=:date, summary=:summary, salary=:salary, overtime=:overtime,
                   meal=:meal, tax=:tax, utilities=:utilities, rent=:rent,
                   municipal=:municipal, pay_method=:pay_method,
                   pay_currency=:pay_currency, total_card=:total_card,
                   total_ves=:total_ves, total_usd=:total_usd,
                   total_cny=:total_cny, notes=:notes
                   WHERE id=:id""",
                {**data, "id": rec_id})
        else:
            cur = conn.execute(
                """INSERT INTO daily_expense
                   (date, summary, salary, overtime, meal, tax, utilities, rent,
                    municipal, pay_method, pay_currency, total_card, total_ves,
                    total_usd, total_cny, notes)
                   VALUES
                   (:date, :summary, :salary, :overtime, :meal, :tax, :utilities,
                    :rent, :municipal, :pay_method, :pay_currency, :total_card,
                    :total_ves, :total_usd, :total_cny, :notes)""",
                data)
            rec_id = cur.lastrowid
        conn.commit()
        return rec_id
    finally:
        conn.close()


def list_daily_expense(date_from=None, date_to=None) -> list:
    conn = get_conn()
    sql = "SELECT * FROM daily_expense"
    conds, args = [], []
    if date_from:
        conds.append("date >= ?")
        args.append(date_from)
    if date_to:
        conds.append("date <= ?")
        args.append(date_to)
    if conds:
        sql += " WHERE " + " AND ".join(conds)
    sql += " ORDER BY date IS NULL, date, id"
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_daily_expense(rec_id: int) -> dict:
    conn = get_conn()
    row = conn.execute("SELECT * FROM daily_expense WHERE id = ?", (rec_id,)).fetchone()
    conn.close()
    return dict(row) if row else None


def delete_daily_expense(rec_id: int):
    conn = get_conn()
    conn.execute("DELETE FROM daily_expense WHERE id = ?", (rec_id,))
    conn.commit()
    conn.close()


# ------------------------------------------------------- daily_expense_items
# 支出日报明细：一笔一行（日期 + 摘要 + 类别 + 付款方式 + 币种 + 金额）

def expense_category_label(key: str) -> str:
    """费用类别 key → 显示名称（内置类别 + 设置里手工新增的自定义类别）。"""
    k = str(key or "").strip()
    if not k:
        return ""
    try:
        categories = settings_mod.get_expense_categories()
    except Exception:  # noqa: BLE001  设置读取失败时回退内置类别
        categories = config.EXPENSE_CATEGORIES
    for ck, label in categories:
        if ck == k:
            return label
    return k


def save_daily_expense_item(rec: dict) -> int:
    conn = get_conn()
    try:
        data = {
            "date": _date_or_iso(rec.get("date")),
            "summary": rec.get("summary", ""),
            "category": rec.get("category", ""),
            "method": rec.get("method", ""),
            "currency": config.normalize_currency(rec.get("currency", "")),
            "amount": float(rec.get("amount") or 0),
            "notes": rec.get("notes", ""),
        }
        rec_id = rec.get("id")
        if rec_id:
            conn.execute(
                """UPDATE daily_expense_items SET
                   date=:date, summary=:summary, category=:category,
                   method=:method, currency=:currency, amount=:amount,
                   notes=:notes WHERE id=:id""",
                {**data, "id": rec_id})
        else:
            cur = conn.execute(
                """INSERT INTO daily_expense_items
                   (date, summary, category, method, currency, amount, notes)
                   VALUES (:date, :summary, :category, :method, :currency,
                           :amount, :notes)""",
                data)
            rec_id = cur.lastrowid
        conn.commit()
        return rec_id
    finally:
        conn.close()


def list_daily_expense_items(date_from=None, date_to=None) -> list:
    conn = get_conn()
    sql = "SELECT * FROM daily_expense_items"
    conds, args = [], []
    if date_from:
        conds.append("date >= ?")
        args.append(date_from)
    if date_to:
        conds.append("date <= ?")
        args.append(date_to)
    if conds:
        sql += " WHERE " + " AND ".join(conds)
    sql += " ORDER BY date IS NULL, date, id"
    try:
        rows = conn.execute(sql, args).fetchall()
    finally:
        conn.close()
    return [dict(r) for r in rows]


def get_daily_expense_item(rec_id: int) -> dict:
    conn = get_conn()
    try:
        row = conn.execute("SELECT * FROM daily_expense_items WHERE id = ?",
                           (rec_id,)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def delete_daily_expense_item(rec_id: int):
    conn = get_conn()
    try:
        conn.execute("DELETE FROM daily_expense_items WHERE id = ?", (rec_id,))
        conn.commit()
    finally:
        conn.close()


# ------------------------------------------------------------------ supplier_settlements

def _calc_supplier_pay(rec: dict):
    """根据方式/币种把支付金额归集到对应列。"""
    amount = float(rec.get("pay_amount", 0) or 0)
    method = rec.get("pay_method", "")
    currency = config.normalize_currency(rec.get("pay_currency", ""))
    bank = amount if method == config.PAY_METHOD_CARD else 0
    cash = method == config.PAY_METHOD_CASH
    cash_ves = amount if (cash and currency == config.PAY_CURRENCY_BS) else 0
    cash_usd = amount if (cash and currency == config.PAY_CURRENCY_USD) else 0
    cash_cny = amount if (cash and currency == config.PAY_CURRENCY_CNY) else 0
    return bank, cash_ves, cash_usd, cash_cny


def save_supplier_settlement(rec: dict) -> int:
    conn = get_conn()
    try:
        bank, cash_ves, cash_usd, cash_cny = _calc_supplier_pay(rec)
        data = {
            "date": _date_or_iso(rec.get("date")),
            "partner_name": rec.get("partner_name", ""),
            "summary": rec.get("summary", ""),
            "doc_number": rec.get("doc_number", ""),
            "supply_amount_ves": rec.get("supply_amount_ves", 0),
            "pay_method": rec.get("pay_method", ""),
            "pay_currency": config.normalize_currency(rec.get("pay_currency", "")),
            "pay_bank": bank,
            "pay_cash_ves": cash_ves,
            "pay_cash_usd": cash_usd,
            "pay_cash_cny": cash_cny,
            "store": rec.get("store", ""),
            "notes": rec.get("notes", ""),
        }
        rec_id = rec.get("id")
        if rec_id:
            conn.execute(
                """UPDATE supplier_settlements SET
                   date=:date, partner_name=:partner_name, summary=:summary,
                   doc_number=:doc_number, supply_amount_ves=:supply_amount_ves,
                   pay_method=:pay_method, pay_currency=:pay_currency,
                   pay_bank=:pay_bank, pay_cash_ves=:pay_cash_ves,
                   pay_cash_usd=:pay_cash_usd, pay_cash_cny=:pay_cash_cny,
                   store=:store, notes=:notes
                   WHERE id=:id""",
                {**data, "id": rec_id})
        else:
            cur = conn.execute(
                """INSERT INTO supplier_settlements
                   (date, partner_name, summary, doc_number, supply_amount_ves,
                    pay_method, pay_currency, pay_bank, pay_cash_ves,
                    pay_cash_usd, pay_cash_cny, store, notes)
                   VALUES
                   (:date, :partner_name, :summary, :doc_number,
                    :supply_amount_ves, :pay_method, :pay_currency, :pay_bank,
                    :pay_cash_ves, :pay_cash_usd, :pay_cash_cny, :store, :notes)""",
                data)
            rec_id = cur.lastrowid
        conn.commit()
        return rec_id
    finally:
        conn.close()


def list_supplier_settlements(date_from=None, date_to=None, partner_name=None) -> list:
    conn = get_conn()
    sql = "SELECT * FROM supplier_settlements"
    conds, args = [], []
    if date_from:
        conds.append("date >= ?")
        args.append(date_from)
    if date_to:
        conds.append("date <= ?")
        args.append(date_to)
    if partner_name:
        conds.append("partner_name LIKE ?")
        args.append(f"%{partner_name}%")
    if conds:
        sql += " WHERE " + " AND ".join(conds)
    sql += " ORDER BY date IS NULL, date, id"
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_supplier_settlement(rec_id: int) -> dict:
    conn = get_conn()
    row = conn.execute("SELECT * FROM supplier_settlements WHERE id = ?",
                       (rec_id,)).fetchone()
    conn.close()
    return dict(row) if row else None


def delete_supplier_settlement(rec_id: int):
    conn = get_conn()
    conn.execute("DELETE FROM supplier_settlements WHERE id = ?", (rec_id,))
    conn.commit()
    conn.close()


# ------------------------------------------------------------------ chart_of_accounts / opening_balances

def list_chart_of_accounts() -> list:
    """返回全部科目。"""
    conn = get_conn()
    rows = conn.execute(
        "SELECT * FROM chart_of_accounts ORDER BY code").fetchall()
    conn.close()
    return [dict(r) for r in rows]


def save_chart_of_account(account: dict):
    """新增/更新单个科目（含明细标志 is_leaf 与上级科目 parent）。"""
    conn = get_conn()
    try:
        is_leaf = account.get("is_leaf")
        is_leaf = 1 if is_leaf is None else (1 if int(is_leaf) else 0)
        parent = account.get("parent")
        if parent is None:
            codes = [r["code"] for r in
                     conn.execute("SELECT code FROM chart_of_accounts").fetchall()]
            parent = infer_parent(account["code"], set(codes) | {account["code"]})
        conn.execute(
            """INSERT INTO chart_of_accounts
               (code, name, direction, is_leaf, parent, is_active)
               VALUES (?,?,?,?,?,?)
               ON CONFLICT(code) DO UPDATE SET
                 name=excluded.name, direction=excluded.direction,
                 is_leaf=excluded.is_leaf, parent=excluded.parent,
                 is_active=excluded.is_active""",
            (account["code"], account.get("name", ""),
             account.get("direction", "借"), is_leaf, parent,
             1 if account.get("is_active", 1) else 0))
        conn.commit()
    finally:
        conn.close()


def delete_chart_of_account(code: str, cascade: bool = True) -> int:
    """删除科目。cascade=True 时连带删除下级科目及其期初余额，返回删除科目数。"""
    conn = get_conn()
    try:
        codes = [code]
        if cascade:
            like = code.replace("%", r"\%") + "%"
            rows = conn.execute(
                "SELECT code FROM chart_of_accounts WHERE code LIKE ? ESCAPE '\\'",
                (like,)).fetchall()
            codes = [r["code"] for r in rows] or [code]
        for c in codes:
            conn.execute("DELETE FROM opening_balances WHERE account_code = ?", (c,))
            conn.execute("DELETE FROM chart_of_accounts WHERE code = ?", (c,))
        conn.commit()
        return len(codes)
    finally:
        conn.close()


def list_opening_balances(year=None) -> list:
    """列出期初余额（含科目名称与层级信息），未录入的科目以 0 补齐。"""
    cols = """SELECT c.code, c.name, c.direction,
                          COALESCE(c.is_leaf,1)         AS is_leaf,
                          COALESCE(c.parent,'')         AS parent,
                          COALESCE(o.opening_balance,0) AS opening_balance,
                          COALESCE(o.cum_debit,0)       AS cum_debit,
                          COALESCE(o.cum_credit,0)      AS cum_credit,
                          COALESCE(o.period_balance,0)  AS period_balance,
                          COALESCE(o.pnl_cum,0)         AS pnl_cum,
                          o.year AS year
                   FROM chart_of_accounts c
                   LEFT JOIN opening_balances o
                     ON o.account_code = c.code"""
    conn = get_conn()
    try:
        if year:
            rows = conn.execute(
                cols + " AND o.year = ? ORDER BY c.code", (year,)).fetchall()
        else:
            rows = conn.execute(cols + " ORDER BY c.code").fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


def save_opening_balances(records: list, year: str):
    """批量保存期初余额（按 account_code + year upsert）。

    records 每项可用 account_code 或 code（list_opening_balances 返回 code）。
    """
    conn = get_conn()
    try:
        for r in records:
            code = r.get("account_code") or r.get("code")
            if not code:
                continue
            conn.execute(
                """INSERT INTO opening_balances
                   (account_code, year, opening_balance, cum_debit,
                    cum_credit, period_balance, pnl_cum)
                   VALUES (?,?,?,?,?,?,?)
                   ON CONFLICT(account_code, year) DO UPDATE SET
                     opening_balance=excluded.opening_balance,
                     cum_debit=excluded.cum_debit,
                     cum_credit=excluded.cum_credit,
                     period_balance=excluded.period_balance,
                     pnl_cum=excluded.pnl_cum""",
                (code, year,
                 float(r.get("opening_balance", 0) or 0),
                 float(r.get("cum_debit", 0) or 0),
                 float(r.get("cum_credit", 0) or 0),
                 float(r.get("period_balance", 0) or 0),
                 float(r.get("pnl_cum", 0) or 0)))
        conn.commit()
    finally:
        conn.close()


def get_opening_balance_total(year=None, only_leaf: bool = False) -> dict:
    """期初余额汇总（用于借贷平衡校验）。

    only_leaf=True 时只统计明细科目（末级），避免汇总科目与下级重复计数。
    """
    where, args = ["1=1"], []
    if year:
        where.append("o.year = ?")
        args.append(year)
    if only_leaf:
        where.append("COALESCE(c.is_leaf, 1) = 1")
    conn = get_conn()
    try:
        rows = conn.execute(
            """SELECT o.account_code, o.period_balance,
                      COALESCE(c.direction, '借') AS direction
               FROM opening_balances o
               LEFT JOIN chart_of_accounts c ON c.code = o.account_code
               WHERE """ + " AND ".join(where), args).fetchall()
        total_debit = 0.0
        total_credit = 0.0
        for r in rows:
            bal = float(r["period_balance"] or 0)
            if r["direction"] == "贷":
                total_credit += bal
            else:
                total_debit += bal
        return {"debit": total_debit, "credit": total_credit,
                "diff": total_debit - total_credit}
    finally:
        conn.close()


def _find_col(col: dict, *keys) -> int:
    """在表头映射 {表头文本: 列号} 中按关键字做包含匹配，返回列号（0 表示未找到）。"""
    for k, c in col.items():
        if all(x in k for x in keys):
            return c
    return 0


def _detect_year(ws, header_row) -> str:
    """从表头之前的说明行（如“启用期间：2026年07期”）提取会计年度。"""
    import re
    for i in range(1, header_row):
        for c in range(1, 12):
            v = ws.cell(row=i, column=c).value
            if v:
                m = re.search(r"(20\d{2})", str(v))
                if m:
                    return m.group(1)
    return ""


def _year_from_filename(path: str) -> str:
    """模板文件名含日期戳（如 20260904..._财务初始余额.xlsx）时兜底取年份。"""
    import re
    m = re.search(r"(20\d{2})", os.path.basename(path or ""))
    return m.group(1) if m else ""


def import_opening_balances_from_xlsx(path: str, year: str = None,
                                      replace: bool = False) -> dict:
    """从《财务初始余额.xlsx》模板导入科目体系与期初余额。

    模板表头（顺序可任意，按名称模糊匹配）：
      *科目编码 | *科目名称 | 明细科目(是/否) | 借贷 |
      年初余额 | 本年累计借方发生额 | 本年累计贷方发生额 |
      期初余额 | 本年累计损益发生额
    科目层级按编码前缀自动推断（100201 → 上级 1002；22210101 → 222101）。

    replace=True：先清空科目表与该年度期初余额，完全以模板重建科目体系。
    返回 {"accounts", "balances", "year", "leaf", "replaced"}
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]

    # 定位表头行：前 8 行内同时含"科目编码"与"科目名称"的行
    header_row, col = None, {}
    for i in range(1, min(8, ws.max_row) + 1):
        vals = {str(ws.cell(row=i, column=c).value or "").strip()
                for c in range(1, 12)}
        if (any(("科目编码" in v or v == "编码") for v in vals)
                and any(("科目名称" in v or v == "名称") for v in vals)):
            header_row = i
            for c in range(1, 12):
                v = str(ws.cell(row=i, column=c).value or "").strip()
                if v:
                    col[v] = c
            break
    if header_row is None:
        raise ValueError(
            "未找到表头行。模板需包含：*科目编码、*科目名称、明细科目(是/否)、借贷")

    c_code = _find_col(col, "科目编码") or _find_col(col, "编码")
    c_name = _find_col(col, "科目名称") or _find_col(col, "名称")
    c_leaf = _find_col(col, "明细科目")
    c_dir = _find_col(col, "借贷") or _find_col(col, "方向")
    c_open = _find_col(col, "年初余额")
    c_debit = _find_col(col, "累计借方")
    c_credit = _find_col(col, "累计贷方")
    c_period = _find_col(col, "期初余额")
    c_pnl = _find_col(col, "损益")
    if not (c_code and c_name):
        raise ValueError("模板缺少“科目编码 / 科目名称”列")

    if year is None:
        year = (_detect_year(ws, header_row)
                or _year_from_filename(path)
                or str(datetime.date.today().year))

    def cell(r, c):
        return ws.cell(row=r, column=c).value if c else 0

    true_vals = ("是", "1", "Y", "y", "TRUE", "True", "true")
    accounts, records = [], []
    for r in range(header_row + 1, ws.max_row + 1):
        code = cell(r, c_code)
        name = cell(r, c_name)
        if code is None or name is None:
            continue
        code = str(code).strip()
        name = str(name).strip()
        if not code or not name:
            continue
        leaf_raw = str(cell(r, c_leaf) or "是").strip()
        direction = str(cell(r, c_dir) or "借").strip()
        accounts.append({
            "code": code, "name": name,
            "direction": direction or "借",
            "is_leaf": 1 if leaf_raw in true_vals else 0})
        records.append({
            "account_code": code,
            "opening_balance": cell(r, c_open) or 0,
            "cum_debit": cell(r, c_debit) or 0,
            "cum_credit": cell(r, c_credit) or 0,
            "period_balance": cell(r, c_period) or 0,
            "pnl_cum": cell(r, c_pnl) or 0,
        })
    if not accounts:
        raise ValueError("模板中未读取到任何科目行")

    codes = [a["code"] for a in accounts]
    conn = get_conn()
    try:
        if replace:
            if year:
                conn.execute("DELETE FROM opening_balances WHERE year = ?", (year,))
            else:
                conn.execute("DELETE FROM opening_balances")
            conn.execute("DELETE FROM chart_of_accounts")
        for a in accounts:
            conn.execute(
                """INSERT INTO chart_of_accounts
                   (code, name, direction, is_leaf, parent, is_active)
                   VALUES (?,?,?,?,?,1)
                   ON CONFLICT(code) DO UPDATE SET
                     name=excluded.name, direction=excluded.direction,
                     is_leaf=excluded.is_leaf, parent=excluded.parent""",
                (a["code"], a["name"], a["direction"], a["is_leaf"],
                 infer_parent(a["code"], codes)))
        for rec in records:
            conn.execute(
                """INSERT INTO opening_balances
                   (account_code, year, opening_balance, cum_debit,
                    cum_credit, period_balance, pnl_cum)
                   VALUES (?,?,?,?,?,?,?)
                   ON CONFLICT(account_code, year) DO UPDATE SET
                     opening_balance=excluded.opening_balance,
                     cum_debit=excluded.cum_debit,
                     cum_credit=excluded.cum_credit,
                     period_balance=excluded.period_balance,
                     pnl_cum=excluded.pnl_cum""",
                (rec["account_code"], year,
                 float(rec["opening_balance"] or 0),
                 float(rec["cum_debit"] or 0),
                 float(rec["cum_credit"] or 0),
                 float(rec["period_balance"] or 0),
                 float(rec["pnl_cum"] or 0)))
        conn.commit()
        if not replace:
            # 为库中原有科目补齐层级关系（已有 is_leaf 的科目保持原值）
            rebuild_account_tree(conn)
    finally:
        conn.close()
    return {"accounts": len(accounts), "balances": len(records), "year": year,
            "leaf": sum(1 for a in accounts if a["is_leaf"]),
            "replaced": bool(replace)}


def list_opening_balance_years() -> list:
    """已录入期初余额的会计年度（降序，去重）。"""
    conn = get_conn()
    try:
        rows = conn.execute(
            """SELECT DISTINCT year FROM opening_balances
               WHERE year IS NOT NULL AND year <> '' ORDER BY year DESC""").fetchall()
        return [r["year"] for r in rows]
    finally:
        conn.close()


# ------------------------------------------------------------------ products / inventory

def save_product(p: dict) -> int:
    """新增/更新商品。p["id"] 存在则更新，否则按 code upsert。返回商品 id。"""
    conn = get_conn()
    try:
        code = (p.get("code") or "").strip()
        pid = p.get("id")
        data = {
            "code": code,
            "barcode": (p.get("barcode") or "").strip(),
            "name": (p.get("name") or "").strip(),
            "category": (p.get("category") or "").strip(),
            "unit": (p.get("unit") or "").strip(),
            "cost_price": float(p.get("cost_price") or 0),
            "sale_price": float(p.get("sale_price") or 0),
            "stock_qty": float(p.get("stock_qty") or 0),
            "min_stock": float(p.get("min_stock") or 0),
            "supplier": (p.get("supplier") or "").strip(),
            "notes": (p.get("notes") or "").strip(),
        }
        if pid:
            conn.execute(
                """UPDATE products SET code=:code, barcode=:barcode, name=:name,
                   category=:category, unit=:unit, cost_price=:cost_price,
                   sale_price=:sale_price, stock_qty=:stock_qty, min_stock=:min_stock,
                   supplier=:supplier, notes=:notes WHERE id=:id""",
                {**data, "id": pid})
        else:
            cur = conn.execute(
                """INSERT INTO products
                   (code, barcode, name, category, unit, cost_price, sale_price,
                    stock_qty, min_stock, supplier, notes)
                   VALUES
                   (:code, :barcode, :name, :category, :unit, :cost_price,
                    :sale_price, :stock_qty, :min_stock, :supplier, :notes)
                   ON CONFLICT(code) DO UPDATE SET
                     barcode=excluded.barcode, name=excluded.name,
                     category=excluded.category, unit=excluded.unit,
                     cost_price=excluded.cost_price, sale_price=excluded.sale_price,
                     stock_qty=excluded.stock_qty, min_stock=excluded.min_stock,
                     supplier=excluded.supplier, notes=excluded.notes""",
                data)
            pid = cur.lastrowid
        conn.commit()
        return pid
    finally:
        conn.close()


def get_product(pid: int) -> dict:
    conn = get_conn()
    try:
        row = conn.execute("SELECT * FROM products WHERE id = ?", (pid,)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def find_product_by_code(code: str) -> dict:
    code = (code or "").strip()
    if not code:
        return None
    conn = get_conn()
    try:
        row = conn.execute(
            "SELECT * FROM products WHERE code = ?", (code,)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def list_products(category: str = None, low_stock: bool = False) -> list:
    conn = get_conn()
    try:
        sql = "SELECT * FROM products"
        conds, args = [], []
        if category:
            conds.append("category = ?")
            args.append(category)
        if low_stock:
            conds.append("min_stock > 0 AND stock_qty < min_stock")
        if conds:
            sql += " WHERE " + " AND ".join(conds)
        sql += " ORDER BY code"
        return [dict(r) for r in conn.execute(sql, args).fetchall()]
    finally:
        conn.close()


def delete_product(pid: int):
    conn = get_conn()
    try:
        conn.execute("DELETE FROM products WHERE id = ?", (pid,))
        conn.commit()
    finally:
        conn.close()


def low_stock_products() -> list:
    """低于库存下限的商品（min_stock>0 且 stock_qty<min_stock）。"""
    return list_products(low_stock=True)


def apply_stock_change(product_id: int, qty: float, move_type: str,
                       ref_doc: int = None, note: str = "", date: str = None) -> float:
    """登记一笔库存变动并更新商品当前库存，返回变动后余额。

    qty 正=入库，负=出库；move_type ∈ compra/venta/ajuste/inventario。
    """
    conn = get_conn()
    try:
        cur = conn.execute(
            "SELECT stock_qty FROM products WHERE id = ?", (product_id,)).fetchone()
        if cur is None:
            return 0.0
        new_balance = float(cur["stock_qty"]) + float(qty)
        conn.execute("UPDATE products SET stock_qty = ? WHERE id = ?",
                     (new_balance, product_id))
        conn.execute(
            """INSERT INTO stock_moves
               (product_id, date, move_type, qty, balance, ref_doc, note)
               VALUES (?,?,?,?,?,?,?)""",
            (product_id, date, move_type, float(qty), new_balance, ref_doc or None,
             note or ""))
        conn.commit()
        return new_balance
    finally:
        conn.close()


def list_stock_moves(limit: int = 100) -> list:
    conn = get_conn()
    try:
        rows = conn.execute(
            """SELECT m.*, p.code AS pcode, p.name AS pname
               FROM stock_moves m LEFT JOIN products p ON p.id = m.product_id
               ORDER BY m.id DESC LIMIT ?""", (limit,)).fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


def rebuild_stock_from_documents() -> dict:
    """以全部单据（compra 入库 / venta 出库，按行项目 code 匹配商品）重算库存。

    清空 stock_moves 并重置每件商品的 stock_qty 为净变动，再按单据顺序重记流水。
    返回 {"products": 命中商品数, "moves": 流水条数}。
    """
    conn = get_conn()
    try:
        conn.execute("DELETE FROM stock_moves")
        conn.execute("UPDATE products SET stock_qty = 0")
        docs = conn.execute(
            "SELECT id, direction, date FROM documents ORDER BY id").fetchall()
        moves = 0
        hit = set()
        for d in docs:
            sign = 1 if d["direction"] == "compra" else -1
            items = conn.execute(
                "SELECT code, qty FROM items WHERE document_id = ?",
                (d["id"],)).fetchall()
            for it in items:
                code = (it["code"] or "").strip()
                if not code:
                    continue
                prod = conn.execute(
                    "SELECT id FROM products WHERE code = ?", (code,)).fetchone()
                if not prod:
                    continue
                qty = float(it["qty"] or 0) * sign
                if qty == 0:
                    continue
                new_balance = float(conn.execute(
                    "SELECT stock_qty FROM products WHERE id = ?",
                    (prod["id"],)).fetchone()["stock_qty"]) + qty
                conn.execute("UPDATE products SET stock_qty = ? WHERE id = ?",
                             (new_balance, prod["id"]))
                conn.execute(
                    """INSERT INTO stock_moves
                       (product_id, date, move_type, qty, balance, ref_doc, note)
                       VALUES (?,?,?,?,?,?,?)""",
                    (prod["id"], d["date"], d["direction"], qty, new_balance,
                     d["id"], "由单据重建"))
                moves += 1
                hit.add(prod["id"])
        conn.commit()
        return {"products": len(hit), "moves": moves}
    finally:
        conn.close()
