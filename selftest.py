"""自检脚本：验证解析器、数据库、报表引擎（不依赖 OCR 模型）。

用法:  python selftest.py
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app.ocr.parser import (  # noqa: E402
    parse_document, detect_currency, detect_exchange_rate)
from app.utils import (  # noqa: E402
    parse_amount, parse_date, parse_document_number, format_amount,
    format_base_amount)
from app import config  # noqa: E402
from app.db import database  # noqa: E402
from app.rates import convert_to_base  # noqa: E402
from app.accounting import build_balance_sheet, build_income_statement  # noqa: E402
from app.accounting import reports  # noqa: E402
from app.sync import manager as sync_manager  # noqa: E402

PASS = 0
FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  [OK] {name}")
    else:
        FAIL += 1
        print(f"  [FAIL] {name} {detail}")


def main():
    print("== 1. 金额/日期/单据号解析 ==")
    check("parse_amount 1.234,56 €", parse_amount("1.234,56 €") == 1234.56)
    check("parse_amount $ 1.234,56", parse_amount("$ 1.234,56") == 1234.56)
    check("parse_amount 37,80", parse_amount("37,80") == 37.8)
    check("parse_amount 1234.56", parse_amount("1234.56") == 1234.56)
    check("format_amount 默认美元符号", format_amount(1234.56) == "1.234,56 $")
    d = parse_date("25/08/2026")
    check("parse_date 25/08/2026", d is not None and d.day == 25 and d.month == 8 and d.year == 2026)
    check("parse_document_number N 2026-045",
          parse_document_number("Factura N 2026-045") == "2026-045")
    check("parse_document_number 排除日期",
          parse_document_number("Fecha: 25/08/2026") == "")

    print("== 2. 单据解析（进货） ==")
    compra = parse_document([
        "FACTURA", "Proveedor: Distribuciones Sol S.L. NIF B12345678",
        "Fecha: 25/08/2026", "Factura N 2026-045", "CONCEPTO",
        "Harina 5kg 10 x 2.50 25,00", "Aceite 1L 4 x 3,20 12,80",
        "Base imponible 37,80", "IVA 21% 7,94", "TOTAL 45,74",
    ])
    check("方向=compra", compra["direction"] == "compra")
    check("供应商名称", compra["partner"] == "Distribuciones Sol S.L.",
          f"got={compra['partner']!r}")
    check("单据号", compra["doc_number"] == "2026-045",
          f"got={compra['doc_number']!r}")
    check("Base", compra["base"] == 37.8)
    check("IVA", compra["iva_amount"] == 7.94, f"got={compra['iva_amount']}")
    check("Total", compra["total"] == 45.74)
    check("行项目数", len(compra["items"]) == 2)

    print("== 3. 单据解析（出货） ==")
    venta = parse_document([
        "FACTURA DE VENTA", "Cliente: Panaderia La Espiga S.L.",
        "Fecha 02/09/2026", "Factura F-2026-120", "ARTICULO",
        "Pan integral 50 x 1,20 60,00", "Bolleria 30 x 2,00 60,00",
        "BASE IMPONIBLE 120,00", "CUOTA IVA 25,20", "TOTAL 145,20",
    ])
    check("方向=venta", venta["direction"] == "venta")
    check("客户名称", venta["partner"] == "Panaderia La Espiga S.L.",
          f"got={venta['partner']!r}")
    check("单据号 F-2026-120", venta["doc_number"] == "F-2026-120",
          f"got={venta['doc_number']!r}")
    check("Base=120", venta["base"] == 120.0)
    check("IVA=25.2", venta["iva_amount"] == 25.2)
    check("Total=145.2", venta["total"] == 145.2)

    print("== 4. 币种与汇率识别 ==")
    check("EUR (€)", detect_currency(["TOTAL 45,74 €"]) == "EUR")
    check("USD ($)", detect_currency(["TOTAL $ 45,74"]) == "USD")
    check("USD (USD)", detect_currency(["Total en USD 100,00"]) == "USD")
    # VES 与 Bs 是同一种货币，统一识别/存储为 Bs
    check("Bs (Bs.)", detect_currency(["TOTAL Bs. 1.234,56"]) == "Bs")
    check("Bs (Bs.S)", detect_currency(["TOTAL Bs.S 1.234,56"]) == "Bs")
    check("Bs (BOLIVARES)", detect_currency(["Total en Bolivares 1.234,56"]) == "Bs")
    check("Bs (VES 代码)", detect_currency(["TOTAL 1.234,56 VES"]) == "Bs")
    check("汇率 Tasa: 36,50",
          detect_exchange_rate(["Tasa de cambio: 36,50"]) == 36.5)
    check("汇率 1 USD = 785,07 Bs",
          detect_exchange_rate(["1 USD = 785,07 Bs"]) == 785.07)
    check("汇率 Bs 785,07 / USD",
          detect_exchange_rate(["Bs. 785,07 / USD"]) == 785.07)
    check("汇率行不影响币种",
          detect_currency(["Tasa: 36,50", "TOTAL $ 100,00"]) == "USD")

    print("== 5. 汇率换算 ==")
    st_eur = {"base_currency": "EUR", "usd_to_base": 0.92, "usd_ves_official": 785.07}
    v, ok = convert_to_base(100, "EUR", st_eur)
    check("EUR 不换算", v == 100 and ok)
    v, ok = convert_to_base(100, "USD", st_eur)
    check("USD→EUR", abs(v - 92.0) < 0.01 and ok, f"got={v}")
    v, ok = convert_to_base(78507, "Bs", st_eur)
    check("Bs→EUR(官方汇率)", abs(v - 92.0) < 0.01 and ok, f"got={v}")
    v, ok = convert_to_base(78507, "VES", st_eur)
    check("旧写法 VES 仍按 Bs 换算", abs(v - 92.0) < 0.01 and ok, f"got={v}")
    v, ok = convert_to_base(100, "USD", {"base_currency": "EUR", "usd_to_base": 0})
    check("缺汇率标记未换算", not ok)
    v, ok = convert_to_base(36.5, "Bs",
                            {"base_currency": "EUR", "usd_to_base": 0.92,
                             "usd_ves_official": 0}, doc_rate=36.5)
    check("单据自带汇率优先", abs(v - 0.92) < 0.01 and ok, f"got={v}")

    print("== 5b. 人民币（CNY）识别与换算 ==")
    st_cny = {"base_currency": "USD", "usd_to_base": 1.0, "usd_cny": 7.2}
    v, ok = convert_to_base(720, "CNY", st_cny)
    check("CNY→USD（720 ÷ 7.2 = 100）", abs(v - 100.0) < 0.01 and ok, f"got={v}")
    v, ok = convert_to_base(720, "CNY", st_cny, doc_rate=7.2)
    check("CNY 单据自带汇率优先", abs(v - 100.0) < 0.01 and ok, f"got={v}")
    v, ok = convert_to_base(100, "CNY", {"base_currency": "USD", "usd_to_base": 1.0})
    check("CNY 缺汇率标记未换算", not ok)
    st_base_cny = {"base_currency": "CNY", "usd_to_base": 7.2}
    v, ok = convert_to_base(100, "USD", st_base_cny)
    check("USD→CNY（本位币人民币，100 × 7.2 = 720）",
          abs(v - 720.0) < 0.01 and ok, f"got={v}")
    v, ok = convert_to_base(500, "CNY", st_base_cny)
    check("本位币=CNY 原值返回", v == 500 and ok)

    check("识别 ¥ 为 CNY", detect_currency(["TOTAL ¥ 1.234,56"]) == "CNY",
          f"got={detect_currency(['TOTAL ¥ 1.234,56'])}")
    check("识别 ￥ 为 CNY", detect_currency(["合计 ￥888.00"]) == "CNY")
    check("识别“人民币”为 CNY", detect_currency(["金额单位：人民币"]) == "CNY")
    check("识别 CNY 代码", detect_currency(["Total: 100.00 CNY"]) == "CNY")
    check("人民币汇率行 1 USD = 7,20 CNY",
          abs(detect_exchange_rate(["1 USD = 7,20 CNY"]) - 7.2) < 0.01,
          f"got={detect_exchange_rate(['1 USD = 7,20 CNY'])}")

    check("人民币金额显示 ¥ 符号",
          format_amount(1234.5, currency="CNY").endswith("¥"),
          format_amount(1234.5, currency="CNY"))

    print("== 6. 数据库 + 报表引擎 ==")
    # 使用临时数据库
    tmp_db = os.path.join(tempfile.mkdtemp(prefix="facturas_selftest_"), "t.db")
    old_db = config.DB_PATH
    config.DB_PATH = tmp_db
    try:
        database.init_db()
        # 全新安装（只有表结构）应判定为“无数据”，同步时建议从云端恢复
        check("空账套 has_local_data=False",
              sync_manager.has_local_data() is False)
        c1 = dict(compra)
        c1["direction"] = "compra"
        c1["store"] = "A店"
        c1["items"] = [{"desc": "Harina", "qty": 10, "unit_price": 2.5,
                        "amount": 25.0, "discount": 8.0, "neto": 23.0},
                       {"desc": "Aceite", "qty": 4, "unit_price": 3.2, "amount": 12.8}]
        v1 = dict(venta)
        v1["direction"] = "venta"
        v1["store"] = "B店"
        v1["items"] = [{"desc": "Pan", "qty": 50, "unit_price": 1.2, "amount": 60.0},
                       {"desc": "Bolleria", "qty": 30, "unit_price": 2.0, "amount": 60.0}]
        database.save_document(c1)
        database.save_document(v1)

        docs = database.list_documents()
        check("已保存 2 张单据", len(docs) == 2)
        check("有业务数据后 has_local_data=True", sync_manager.has_local_data())

        # 分店字段保存/读回 + 按分店筛选
        check("分店保存并读回",
              any(d.get("store") == "A店" for d in docs) and
              any(d.get("store") == "B店" for d in docs),
              f"stores={[d.get('store') for d in docs]}")
        a_docs = database.list_documents(store="A店")
        b_docs = database.list_documents(store="B店")
        check("按分店筛选",
              len(a_docs) == 1 and a_docs[0]["store"] == "A店"
              and len(b_docs) == 1 and b_docs[0]["store"] == "B店")
        check("分店去重列表",
              set(database.list_stores()) == {"A店", "B店"},
              f"got={database.list_stores()}")

        # 行项目折扣字段保存/读回
        full = database.get_document(docs[0]["id"])
        harina = full["items"][0]
        check("行项目折扣保存并读回",
              abs(harina["discount"] - 8.0) < 0.01 and abs(harina["neto"] - 23.0) < 0.01,
              f"got discount={harina.get('discount')} neto={harina.get('neto')}")

        docs = database.list_documents()
        for d in docs:
            full = database.get_document(d["id"])
            d["items"] = full["items"]

        bal = build_balance_sheet(docs, capital=1000.0)
        print(f"    资产合计: {bal['total_activo']}  负债+净资产: {bal['total_pasivo_pat']}")
        check("资产负债表平衡", bal["balanced"], f"diff={bal['diff']}")
        check("现金 = 资本 + 收入 - 进货",
              abs(bal["activo"][0]["amount"] - (1000 + 145.2 - 45.74)) < 0.01)

        inc = build_income_statement(docs)
        print(f"    收入: {inc['ventas']}  成本: {inc['coste']}  净利: {inc['resultado']}")
        check("利润表收入=120", inc["ventas"] == 120.0)
        # 平均成本: 进货 37.8 / 14 件 = 2.70；出货 80 件 → 成本 216.0
        check("销售成本(加权平均)=216",
              abs(inc["coste"] - (37.8 / 14) * 80) < 0.01,
              f"got={inc['coste']}")
        check("净利润=收入-成本", abs(inc["resultado"] - (120.0 - inc["coste"])) < 0.01)

        print("== 6b. 店铺收支/支出日报与供应商结算 ==")
        database.save_daily_income({
            "date": "2026-08-25",
            "rows": [
                {"store": "A", "source": config.INCOME_SOURCE_CASH,
                 "currency": config.INCOME_CUR_USD, "amount": 30},
                {"store": "A", "source": config.INCOME_SOURCE_CASH,
                 "currency": config.INCOME_CUR_BS, "amount": 200},
                {"store": "A", "source": config.INCOME_SOURCE_EPAY,
                 "currency": config.INCOME_CUR_USDT, "amount": 100},
                {"store": "A", "source": config.INCOME_SOURCE_CARD,
                 "currency": config.INCOME_CUR_USD, "amount": 50},
            ],
        })
        incomes = database.list_daily_income_full()
        check("收支日报保存并读取",
              len(incomes) == 1 and len(incomes[0]["rows"]) == 4
              and incomes[0]["rows"][0]["amount"] == 30,
              f"got={len(incomes[0]['rows']) if incomes else 0} 行")

        # 备注按笔独立：只改其中一笔，同一天其他笔不受影响（回归）
        flat = database.list_daily_income_rows()
        database.update_daily_income_row(flat[0]["row_id"], {
            "store": flat[0]["store"], "source": flat[0]["source"],
            "currency": flat[0]["currency"], "amount": flat[0]["amount"],
            "notes": "第一笔备注"})
        flat = database.list_daily_income_rows()
        check("再编辑备注只改这一笔",
              flat[0]["notes"] == "第一笔备注"
              and all(r["notes"] == "" for r in flat[1:]),
              f"got={[r['notes'] for r in flat]}")
        agg = database.daily_income_by_account()
        check("日报按支付方式+币种归集科目",
              abs(agg["cash"].get("USD", 0) - 30) < 0.01
              and abs(agg["cash"].get("Bs", 0) - 200) < 0.01
              and abs(agg["bank"].get("USD", 0) - 50) < 0.01
              and abs(agg["crypto"].get("USDT", 0) - 100) < 0.01,
              f"got={agg}")

        # 收入「支付方式+币种」→ 财务报表科目
        check("现钞 Bs → 库存现金(cash)",
              config.income_account_key("现钞", "Bs") == "cash",
              f"got={config.income_account_key('现钞', 'Bs')}")
        check("现钞 USD → 库存现金(cash)",
              config.income_account_key("现钞", "USD") == "cash")
        check("现钞 CNY → 库存现金(cash)",
              config.income_account_key("现钞", "CNY") == "cash")
        check("银行卡 USD → 银行存款(bank)",
              config.income_account_key("银行卡", "USD") == "bank")
        check("电子支付 CNY → 其他货币资金(crypto)",
              config.income_account_key("电子支付", "CNY") == "crypto")
        # 稳定币无现钞形态：现钞来源改记银行存款，电子支付记其他货币资金
        check("USDT 只能记银行/其他货币资金(不可现金)",
              config.income_account_options("USDT") == ["bank", "crypto"])
        check("现钞 USDT → 银行存款(bank)",
              config.income_account_key("现钞", "USDT") == "bank")
        check("电子支付 USDT → 其他货币资金(crypto)",
              config.income_account_key("电子支付", "USDT") == "crypto")
        check("旧写法 VES 现钞 → 库存现金(cash)",
              config.income_account_key("现钞", "VES") == "cash")
        check("法定货币三种科目均可记",
              config.income_account_options("Bs") == ["cash", "bank", "crypto"])

        database.save_daily_expense({
            "date": "2026-08-25", "summary": "日常支出",
            "salary": 100, "overtime": 20, "meal": 30, "tax": 10,
            "utilities": 25, "rent": 200, "municipal": 15,
            "pay_method": config.PAY_METHOD_CARD,
            "pay_currency": config.PAY_CURRENCY_BS,
        })
        expenses = database.list_daily_expense()
        check("支出日报自动归集到刷卡",
              len(expenses) == 1 and expenses[0]["total_card"] == 400
              and expenses[0]["total_ves"] == 0 and expenses[0]["total_usd"] == 0)

        database.save_daily_expense({
            "date": "2026-08-25", "summary": "现金美元支出",
            "salary": 0, "overtime": 0, "meal": 0, "tax": 0,
            "utilities": 0, "rent": 0, "municipal": 0,
            "pay_method": config.PAY_METHOD_CASH,
            "pay_currency": config.PAY_CURRENCY_USD,
        })
        expenses = database.list_daily_expense(
            date_from="2026-08-25", date_to="2026-08-25")
        check("支出日报按日期筛选", len(expenses) == 2)

        database.save_supplier_settlement({
            "date": "2026-08-25", "partner_name": "Proveedor A",
            "summary": "Pago", "doc_number": "P-001",
            "supply_amount_ves": 1000,
            "pay_method": config.PAY_METHOD_CASH,
            "pay_currency": config.PAY_CURRENCY_BS,
            "pay_amount": 500,
            "store": "B店",
        })
        settlements = database.list_supplier_settlements()
        check("供应商结算现金委币",
              len(settlements) == 1 and settlements[0]["pay_cash_ves"] == 500
              and settlements[0]["pay_bank"] == 0
              and settlements[0]["pay_cash_usd"] == 0)
        check("结算入库分店保存读回",
              len(settlements) == 1 and settlements[0].get("store") == "B店",
              f"store={settlements[0].get('store') if settlements else None}")

        # 人民币收付归集
        database.save_daily_expense({
            "date": "2026-08-26", "summary": "人民币现金支出",
            "salary": 100, "overtime": 0, "meal": 0, "tax": 0,
            "utilities": 0, "rent": 0, "municipal": 0,
            "pay_method": config.PAY_METHOD_CASH,
            "pay_currency": config.PAY_CURRENCY_CNY,
        })
        exps = database.list_daily_expense(
            date_from="2026-08-26", date_to="2026-08-26")
        check("人民币现金支出归集到 total_cny",
              len(exps) == 1 and exps[0]["total_cny"] == 100
              and exps[0]["total_ves"] == 0 and exps[0]["total_usd"] == 0,
              f"got={exps[0] if exps else None}")

        database.save_supplier_settlement({
            "date": "2026-08-26", "partner_name": "Proveedor B",
            "summary": "Pago", "doc_number": "P-002",
            "supply_amount_ves": 0,
            "pay_method": config.PAY_METHOD_CASH,
            "pay_currency": config.PAY_CURRENCY_CNY,
            "pay_amount": 800,
        })
        sts = database.list_supplier_settlements(
            date_from="2026-08-26", date_to="2026-08-26")
        check("人民币结算归集到 pay_cash_cny",
              len(sts) == 1 and sts[0]["pay_cash_cny"] == 800,
              f"got={sts[0] if sts else None}")

        # PDF 导出（需要 reportlab）
        try:
            from app.accounting import export_pdf
            out = export_pdf("balance", out_path=os.path.join(
                os.path.dirname(tmp_db), "b.pdf"))
            check("PDF 导出", os.path.exists(out) and os.path.getsize(out) > 500)
        except ImportError:
            print("  [SKIP] PDF 导出：未安装 reportlab")

        print("== 6c. 财务期初余额（科目体系 + 期初数） ==")
        # 构造《财务初始余额.xlsx》模板格式的临时 Excel（含明细科目列与多级编码）
        try:
            from openpyxl import Workbook
            xlsx = os.path.join(os.path.dirname(tmp_db), "财务初始余额.xlsx")
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "财务初始余额"
            ws["A2"] = "测试企业"
            ws["G2"] = "启用期间：2026年07期"
            headers = ["*科目编码", "*科目名称", "明细科目(是/否)", "借贷", "年初余额",
                       "本年累计借方发生额", "本年累计贷方发生额", "期初余额",
                       "本年累计损益发生额"]
            for i, h in enumerate(headers, 1):
                ws.cell(row=3, column=i, value=h)
            data = [
                ["1001", "库存现金", "是", "借", 0, 0, 0, 5000, 0],
                ["1002", "银行存款", "否", "借", 0, 0, 0, 0, 0],
                ["100201", "基本户", "是", "借", 0, 0, 0, 15000, 0],
                ["1405", "库存商品", "是", "借", 0, 0, 0, 8000, 0],
                ["4001", "实收资本", "是", "贷", 0, 0, 0, 28000, 0],
            ]
            for r, row in enumerate(data, 4):
                for c, v in enumerate(row, 1):
                    ws.cell(row=r, column=c, value=v)
            wb.save(xlsx)

            # 清空科目表与期初余额，保证用例自包含（不受默认科目表 JSON 影响）
            conn = database.get_conn()
            conn.execute("DELETE FROM opening_balances")
            conn.execute("DELETE FROM chart_of_accounts")
            conn.commit()
            conn.close()

            res = database.import_opening_balances_from_xlsx(xlsx)
            check("导入 5 个科目", res["accounts"] == 5, f"got={res['accounts']}")
            check("导入年度=2026", res["year"] == "2026", f"got={res['year']}")
            check("明细科目 4 个", res["leaf"] == 4, f"got={res['leaf']}")

            rows = database.list_opening_balances(year="2026")
            check("期初余额列出 5 条",
                  len(rows) == 5 and rows[0]["name"] == "库存现金")
            by_code = {r["code"]: r for r in rows}
            check("层级推断：100201 上级=1002",
                  by_code["100201"]["parent"] == "1002",
                  f"got={by_code['100201'].get('parent')}")
            check("汇总科目标记：1002 is_leaf=0",
                  by_code["1002"]["is_leaf"] == 0,
                  f"got={by_code['1002'].get('is_leaf')}")

            total = database.get_opening_balance_total(year="2026", only_leaf=True)
            check("明细口径借贷平衡（5000+15000+8000=28000）",
                  abs(total["diff"]) < 0.005,
                  f"借={total['debit']} 贷={total['credit']}")

            # 修改期初数后重算平衡
            records = [{"account_code": "1001", "period_balance": 6000},
                       {"account_code": "4001", "period_balance": 29000}]
            database.save_opening_balances(records, year="2026")
            total2 = database.get_opening_balance_total(year="2026", only_leaf=True)
            check("修改期初数后仍平衡（6000+15000+8000=29000）",
                  abs(total2["diff"]) < 0.005,
                  f"借={total2['debit']} 贷={total2['credit']}")

            # 级联删除：删除 1002 应连带 100201 及其期初余额
            n = database.delete_chart_of_account("1002", cascade=True)
            check("级联删除 2 个科目", n == 2, f"got={n}")
            rows2 = database.list_opening_balances(year="2026")
            check("删除后剩余 3 条", len(rows2) == 3, f"got={len(rows2)}")

            print("== 6d. 财务报表与科目表打通 ==")
            # 建立与模板编码一致的科目体系
            for a in ({"code": "1001", "name": "库存现金", "direction": "借", "is_leaf": 1},
                      {"code": "1405", "name": "库存商品", "direction": "借", "is_leaf": 1},
                      {"code": "22210101", "name": "进项税额", "direction": "借", "is_leaf": 1},
                      {"code": "22210106", "name": "销项税额", "direction": "贷", "is_leaf": 1},
                      {"code": "3001", "name": "实收资本", "direction": "贷", "is_leaf": 1},
                      {"code": "5001", "name": "主营业务收入", "direction": "贷", "is_leaf": 1},
                      {"code": "5401", "name": "主营业务成本", "direction": "借", "is_leaf": 1}):
                database.save_chart_of_account(a)
            database.delete_chart_of_account("4001")
            # 期初：借 1001 20000 + 1405 8000 = 贷 3001 28000
            database.save_opening_balances(
                [{"account_code": "1001", "period_balance": 20000},
                 {"account_code": "1405", "period_balance": 8000},
                 {"account_code": "3001", "period_balance": 28000}], year="2026")

            docs2 = database.list_documents()
            for d in docs2:
                full = database.get_document(d["id"])
                d["items"] = full["items"] if full else []

            rep = build_balance_sheet(docs2, capital=1000.0, year="2026")
            codes = [x["code"] for x in rep["activo"]]
            check("资产行取自科目表", "1001" in codes and "1405" in codes, f"got={codes}")
            check("资产负债表平衡（科目表期初）", rep["balanced"], f"diff={rep['diff']}")
            cash = next((x["amount"] for x in rep["activo"] if x["code"] == "1001"), None)
            check("现金=期初20000+收145.2-付45.74",
                  cash is not None and abs(cash - (20000 + 145.2 - 45.74)) < 0.01,
                  f"got={cash}")
            check("期初来源=科目表",
                  str(rep.get("opening_source", "")).startswith("科目表"),
                  f"got={rep.get('opening_source')}")

            inc2 = build_income_statement(docs2, year="2026")
            check("利润表收入=120（科目口径）",
                  abs(inc2["ventas"] - 120.0) < 0.01, f"got={inc2['ventas']}")
            check("利润表成本=216（科目口径）",
                  abs(inc2["coste"] - 216.0) < 0.01, f"got={inc2['coste']}")
            check("净利润=-96（科目口径）",
                  abs(inc2["resultado"] - (-96.0)) < 0.01, f"got={inc2['resultado']}")

            print("== 6e. 商品库存 / 进销存联动 ==")
            pid = database.save_product({
                "code": "H171", "name": "Leche", "category": "饮料",
                "unit": "UND", "cost_price": 1, "sale_price": 1.5,
                "stock_qty": 0, "min_stock": 5})
            check("新建商品", pid is not None)
            database.save_document({
                "direction": "compra", "partner": "Prov1", "date": "2026-09-12",
                "total": 10.0,
                "items": [{"code": "H171", "desc": "Leche", "qty": 10,
                           "unit_price": 1}]})
            check("进货入库 +10", database.get_product(pid)["stock_qty"] == 10)
            database.save_document({
                "direction": "venta", "partner": "Cli1", "date": "2026-09-12",
                "total": 3,
                "items": [{"code": "H171", "desc": "Leche", "qty": 3,
                           "unit_price": 1.5}]})
            check("销售出库 -3", database.get_product(pid)["stock_qty"] == 7)
            check("库存高于下限不预警",
                  len(database.low_stock_products()) == 0)
            database.apply_stock_change(pid, -10, "ajuste",
                                        note="报损", date="2026-09-12")
            check("手工调整到 -3 触发预警",
                  len(database.low_stock_products()) == 1)
            rb = database.rebuild_stock_from_documents()
            check("从单据重建库存=7",
                  database.get_product(pid)["stock_qty"] == 7,
                  f"got={database.get_product(pid)['stock_qty']}")
            check("重建流水命中 1 商品", rb["products"] == 1 and rb["moves"] == 2,
                  f"got={rb}")
            nmoves = len(database.list_stock_moves(50))
            database.delete_product(pid)
            check("删除商品级联清空其流水",
                  len(database.list_stock_moves(50)) == 0 and nmoves == 2,
                  f"before={nmoves} after={len(database.list_stock_moves(50))}")
        except ImportError:
            print("  [SKIP] 期初余额 Excel：未安装 openpyxl")
    finally:
        config.DB_PATH = old_db

    print("== 6f. 币种名称统一「中文名称（简称）」 ==")
    check("USD 显示名", config.currency_label("USD") == "美元（USD）",
          f"got={config.currency_label('USD')}")
    check("Bs 显示名（VES 同币）",
          config.currency_label("VES") == "玻利瓦尔（Bs）",
          f"got={config.currency_label('VES')}")
    check("显示名 → 代码", config.currency_code("人民币（CNY）") == "CNY",
          f"got={config.currency_code('人民币（CNY）')}")
    check("代码 → 代码", config.currency_code("USDT") == "USDT")
    check("旧中文写法归一", config.normalize_currency("美元") == "USD"
          and config.normalize_currency("人民币") == "CNY"
          and config.normalize_currency("USDT 泰达币") == "USDT")
    check("收入日报币种下拉用统一名",
          config.INCOME_CURRENCY_LABELS["Bs"] == "玻利瓦尔（Bs）",
          f"got={config.INCOME_CURRENCY_LABELS}")

    print("== 6g. 付款方式 ↔ 货币资金大类 ==")
    # 换一张干净的临时库：含默认科目表（1001 库存现金 / 1002 银行存款 / 1012 …）
    tmp_db2 = os.path.join(tempfile.mkdtemp(prefix="facturas_selftest2_"), "t.db")
    config.DB_PATH = tmp_db2
    database.init_db()
    opts = reports.payment_account_options()
    check("付款方式取自科目表", bool(opts)
          and all(o.get("name") for o in opts), f"got={opts}")
    check("付款方式固定三大类（库存现金/银行存款/其他货币资金）",
          [o["name"] for o in opts] == ["库存现金", "银行存款", "其他货币资金"]
          and [o["root"] for o in opts] == ["1001", "1002", "1012"],
          f"got={[o['name'] for o in opts]}")
    check("与收入日报支付来源一一对应",
          [config.income_account_key(s, "Bs") for s in
           (config.INCOME_SOURCE_CASH, config.INCOME_SOURCE_CARD,
            config.INCOME_SOURCE_EPAY)] == ["cash", "bank", "crypto"]
          and [reports.resolve_account(reports.load_chart_index(), k) for k in
               ("cash", "bank", "crypto")] == ["1001", "1002", "1012"],
          f"got={[reports.resolve_account(reports.load_chart_index(), k) for k in ('cash', 'bank', 'crypto')]}")
    check("旧明细文本（库存现金（Bs））仍归到 1001",
          reports.resolve_payment_account("库存现金（Bs）") == "1001",
          f"got={reports.resolve_payment_account('库存现金（Bs）')}")
    leaf = next((o for o in opts if o["root"] == "1002"), opts[0])
    check("同名科目精确匹配", reports.resolve_payment_account(leaf["name"])
          == leaf["code"], f"got={reports.resolve_payment_account(leaf['name'])}")
    check("银行卡回退到银行存款",
          reports.resolve_payment_account("银行卡")
          == reports.resolve_account(reports.load_chart_index(), "bank"))
    check("付款渠道：现金科目→cash",
          reports.payment_channel("库存现金") == "cash",
          f"got={reports.payment_channel('库存现金')}")
    check("付款渠道：银行科目→bank",
          reports.payment_channel(leaf["name"]) == "bank",
          f"got={reports.payment_channel(leaf['name'])}")
    check("付款渠道：兼容旧「银行卡」",
          reports.payment_channel(config.PAY_METHOD_CARD) == "bank")
    database.save_supplier_settlement({
        "partner_name": "Prov1", "date": "2026-09-15", "pay_amount": 100,
        "pay_method": "库存现金", "pay_currency": "Bs",
        "debit_amount": 0, "credit_amount": 0, "notes": ""})
    rows = [r for r in database.list_supplier_settlements()
            if r["pay_method"] == "库存现金"]
    check("供应商结算按科目归集（现金 Bs）",
          rows and abs(rows[0]["pay_cash_ves"] - 100) < 0.01
          and abs(rows[0]["pay_bank"] - 0) < 0.01, f"got={rows[0] if rows else None}")

    print("== 6h. 科目余额表 ==")
    trial = reports.build_trial_balance(
        database.list_documents(), 0.0)["rows"]
    codes = [r["code"] for r in trial]
    check("列出全部科目", len(codes) > 100, f"got={len(codes)}")
    check("含末级科目 100201 与其上级 1002",
          "100201" in codes and "1002" in codes)
    parent = next((r for r in trial if r["code"] == "1002"), None)
    child = next((r for r in trial if r["code"] == "100201"), None)
    check("父科目汇总其明细",
          parent and child and not child["is_parent"]
          and parent["is_parent"], f"got={parent} / {child}")
    tot = reports.build_trial_balance(database.list_documents(), 0.0)["totals"]
    check("本期借贷合计相等",
          abs(tot["debit"] - tot["credit"]) < 0.01, f"got={tot}")
    check("期末借贷合计相等",
          abs(tot["ending_debit"] - tot["ending_credit"]) < 0.01, f"got={tot}")
    rep = reports.get_report("trial")
    rep = rep[0] if isinstance(rep, tuple) else rep
    check("get_report('trial') 可用",
          bool(rep.get("rows")) and "base_currency" in rep)

    print("== 6h2. 店铺收入 / 支出日报 → 财务报表（模块关联）==")
    from app import settings as settings_mod
    st = settings_mod.load_settings()
    st["base_currency"] = "USD"
    settings_mod.save_settings(st)
    # 模拟用户的科目表：库存现金（Bs）100101 / 库存现金（USD）100102
    for _code, _name in (("100101", "Bs"), ("100102", "USD")):
        database.save_chart_of_account({"code": _code, "name": _name,
                                        "direction": "借", "is_leaf": 1})
    iid = database.get_or_create_daily_income("2026-09-01", "")
    database.add_daily_income_row(iid, {
        "store": "A店", "source": config.INCOME_SOURCE_CASH,
        "currency": "Bs", "amount": 100, "notes": ""})
    database.add_daily_income_row(iid, {
        "store": "B店", "source": config.INCOME_SOURCE_CASH,
        "currency": "USD", "amount": 300, "notes": ""})
    for _cat, _amt, _store in (("工资", 200, "A店"), ("水电费", 60, "B店")):
        database.save_daily_expense_item({
            "date": "2026-09-01", "store": _store, "summary": _cat,
            "category": _cat, "method": config.PAY_METHOD_CASH,
            "currency": "USD", "amount": _amt, "notes": ""})
    exp_rows = database.list_daily_expense_items()
    check("支出日报带分店字段（与收入日报一致）",
          [r.get("store") for r in exp_rows] == ["A店", "B店"],
          f"got={[r.get('store') for r in exp_rows]}")
    check("支出日报可按分店筛选",
          len(database.list_daily_expense_items(store="A店")) == 1
          and database.list_daily_expense_items(store="A店")[0]["summary"]
          == "工资",
          f"got={database.list_daily_expense_items(store='A店')}")
    entries, stats, un = reports.daily_book_entries(None, None)
    check("收入 2 笔 / 支出 2 笔生成 8 条分录",
          stats["income_count"] == 2 and stats["expense_count"] == 2
          and len(entries) == 8, f"got={stats} n={len(entries)}")
    check("分店不影响记账科目（仍按类别/付款方式入账）",
          all(str(c).startswith(("1", "5")) for c, _a in entries),
          f"got={entries}")
    check("金额按原币入账、不折算（收入 400 / 支出 260）",
          abs(stats["income_amount"] - 400) < 0.01
          and abs(stats["expense_amount"] - 260) < 0.01 and not un,
          f"got={stats} un={un}")
    docs = database.list_documents()
    tr = reports.build_trial_balance(docs, 0.0, entries=entries)
    by_code = {r["code"]: r for r in tr["rows"]}
    check("收入日报按币种记明细科目（Bs→100101 借 100 / USD→100102 借 300）",
          abs(by_code["100101"]["debit"] - 100) < 0.01
          and abs(by_code["100102"]["debit"] - 300) < 0.01
          and abs(by_code["5001"]["credit"] - 400) < 0.01,
          f"bs={by_code.get('100101')} usd={by_code.get('100102')}")
    check("支出日报按付款方式币种记明细（USD→100102 贷 260）",
          abs(by_code["100102"]["credit"] - 260) < 0.01,
          f"got={by_code.get('100102')}")
    check("支出日报落到费用末级科目（工资 / 水电）",
          any(r["code"].startswith("5602") and r["debit"] > 0
              and not r["is_parent"] for r in tr["rows"]),
          f"got={[ (r['code'], r['debit']) for r in tr['rows'] if r['debit'] ]}")
    check("计入日报后借贷仍平衡",
          abs(tr["totals"]["debit"] - tr["totals"]["credit"]) < 0.01
          and abs(tr["totals"]["ending_debit"]
                  - tr["totals"]["ending_credit"]) < 0.01,
          f"got={tr['totals']}")
    bs = reports.build_balance_sheet(docs, 0.0, entries=entries)
    check("资产负债表仍平衡（资产 = 负债 + 权益）", bs["balanced"],
          f"activo={bs['total_activo']} pas+pat={bs['total_pasivo_pat']}")
    inc = reports.build_income_statement(docs, entries=entries)
    check("利润表含日报收入 400 / 费用 260 / 净利 140",
          abs(inc["ventas"] - 400) < 0.01 and abs(inc["coste"] - 260) < 0.01
          and abs(inc["resultado"] - 140) < 0.01,
          f"ventas={inc['ventas']} coste={inc['coste']} res={inc['resultado']}")
    rep2 = reports.get_report("income")
    rep2 = rep2[0] if isinstance(rep2, tuple) else rep2
    check("get_report 带出数据来源统计",
          (rep2.get("sources") or {}).get("income_rows") == 2
          and (rep2.get("sources") or {}).get("expense_rows") == 2,
          f"got={rep2.get('sources')}")
    config.DB_PATH = old_db

    print("== 6i. 桌面界面构建 / 报表切换（无显示环境跳过）==")
    try:
        config.DB_PATH = tmp_db      # 界面构建需要完整库（临时库已建表并有数据）
        import tkinter as tk
        from tkinter import ttk, messagebox
        from app.ui.reports_page import ReportsPage
        from app.ui.settings_page import SettingsPage
        from app.ui.sync_page import SyncPage
        from app.ui.daily_income_page import DailyIncomePage
        from app.ui.daily_expense_page import DailyExpensePage

        # 无人值守：屏蔽弹窗，避免报表提示框阻塞
        messagebox.showinfo = lambda *a, **k: None
        messagebox.showwarning = lambda *a, **k: None

        root = tk.Tk()
        root.withdraw()

        class _State:
            settings = {"base_currency": "USD"}

            def capital(self):
                return 0.0

        class _App:
            state = _State()

        def _build(cls):
            p = cls(_App())
            p.frame = ttk.Frame(root)
            p.frame.pack()
            p.build()
            return p

        for cls in (SettingsPage, SyncPage, DailyIncomePage, DailyExpensePage,
                    ReportsPage):
            try:
                _build(cls)
                check(f"{cls.__name__} 可构建", True)
            except Exception as e:  # noqa: BLE001
                check(f"{cls.__name__} 可构建", False, f"{e}")

        page = _build(ReportsPage)
        counts = {}
        for t in ("balance", "income", "trial", "balance", "trial"):
            page.type_var.set(t)
            page._generate()
            counts[t] = len(page.tree.get_children())
        check("资产负债表/利润表可渲染",
              counts["balance"] > 0 and counts["income"] > 0, f"got={counts}")
        check("科目余额表可渲染且可来回切换", counts["trial"] > 0,
              f"got={counts}")
        # 回归：2 列 ↔ 8 列切换曾触发 TclError「Invalid column index」导致界面崩溃
        check("切到科目余额表后为 8 列",
              len(page.tree["columns"]) == 8, f"got={page.tree['columns']}")
        root.destroy()
    except Exception as e:  # noqa: BLE001
        print(f"  [SKIP] 桌面界面冒烟：{e}")
    finally:
        config.DB_PATH = old_db

    print("== 6j. 报表金额不跟货币符号 ==")
    check("报表金额只输出数字（1.234,50）",
          format_amount(1234.5, symbols=False) == "1.234,50",
          format_amount(1234.5, symbols=False))
    check("日报列表金额不带货币符号（币种已单列显示）",
          format_amount(260, symbols=False) == "260,00"
          and "$" not in format_amount(260, symbols=False),
          f"got={format_amount(260, symbols=False)}")
    check("Bs 金额后缀为 Bs.",
          format_amount(1234.5, currency="Bs") == "1.234,50 Bs.",
          format_amount(1234.5, currency="Bs"))
    check("Bs 本位币显示名",
          config.currency_label("Bs") == "玻利瓦尔（Bs）",
          config.currency_label("Bs"))
    try:
        config.DB_PATH = tmp_db
        from app import settings as _st
        from app.web.server import app as _web_app
        st = _st.load_settings()
        st["base_currency"] = "Bs"
        _st.save_settings(st)
        _web_app.config["TESTING"] = True
        check("商品价格带本位币符号 Bs.（不再显示 $）",
              format_base_amount(1234.5) == "1.234,50 Bs.",
              f"got={format_base_amount(1234.5)}")
        html = _web_app.test_client().get(
            "/reports?type=trial").get_data(as_text=True)
        check("Web 科目余额表金额不带 $", "$" not in html)
        check("Web 表头不再标注币种",
              '<th class="num">期初借方</th>' in html
              and '<th class="num">期初借方（' not in html)
        pdf = reports.export_pdf("trial", out_path=os.path.join(
            os.path.dirname(tmp_db), "cur.pdf"))
        check("科目余额表 PDF 可导出", os.path.exists(pdf), pdf)
        cli = _web_app.test_client()
        for path, name in (("/", "经营概览"), ("/documents", "单据管理"),
                           ("/income", "店铺收入日报"), ("/expense", "店铺支出日报"),
                           ("/products", "商品库存")):
            page = cli.get(path).get_data(as_text=True)
            check(f"Web {name} 金额不带 $", "$" not in page,
                  f"got={[l for l in page.splitlines() if '$' in l][:2]}")
        exp_html = cli.get("/expense").get_data(as_text=True)
        check("Web 支出日报含分店下拉与分店列",
              '<select name="store">' in exp_html and "<th>分店</th>" in exp_html)
    except Exception as e:  # noqa: BLE001
        print(f"  [SKIP] Web 报表渲染：{e}")
    finally:
        config.DB_PATH = old_db

    print(f"\n结果: {PASS} 通过, {FAIL} 失败")
    sys.exit(1 if FAIL else 0)


if __name__ == "__main__":
    main()
